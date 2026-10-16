"""Toplu QR içe aktarma: dosya okuma (CSV/XLSX), sütun eşleştirme ve satır doğrulama.

Bu modül veritabanına YAZMAZ (yalnızca okur). Kayıt işlemi panel görünümünde,
tek bir transaction içinde servisler.qr_kaydet ile yapılır.
Sunucu dosyadaki URL'leri ASLA çağırmaz; yalnızca biçimini doğrular.
"""
import csv
import io
import posixpath
import zipfile
from xml.parsers import expat  # yalnızca akış tarayıcısı: DOCTYPE/ENTITY işleyicileri istisna fırlatır (varlık genişletme yok)
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from zoneinfo import ZoneInfo

import defusedxml.ElementTree as ET  # XXE / billion laughs / harici varlık: reddedilir (DefusedXmlException)
from defusedxml import DefusedXmlException
from django.core.exceptions import ValidationError

from .dogrulama import guvenli_url_dogrula
from .models import (
    GenelAyarlar, Kampanya, KampanyaSayfasi, Magaza, QRKod, Urun, ayrilmis_kod_hatasi, kisa_kod_dogrula, otomatik_kisa_kod,
)
from .servisler import kod_uretici

EN_FAZLA_BAYT = 5 * 1024 * 1024
EN_FAZLA_SATIR = 20000
EN_FAZLA_SUTUN = 60
EN_FAZLA_BOS_SATIR = 1000  # boş satırlar dahil okunabilecek toplam satır: EN_FAZLA_SATIR + EN_FAZLA_BOS_SATIR
EN_FAZLA_ACILMIS_BAYT = 10 * 1024 * 1024  # xlsx (zip) içeriğinin açılmış toplam boyutu
ZAMAN_DILIMI = ZoneInfo("Europe/Istanbul")

ALANLAR = [
    ("kod", "Kısa kod"),
    ("ad", "QR kod adı"),
    ("baslik", "Hedef başlığı"),
    ("tur", "Tür"),
    ("hedef_url", "Hedef URL"),
    ("sayfa", "Hedef sayfa (slug)"),
    ("baslangic", "Başlangıç"),
    ("bitis", "Bitiş"),
    ("durum", "Durum"),
    ("urun_kodu", "Ürün kodu"),
    ("urun_adi", "Ürün adı"),
    ("kampanya_kodu", "Kampanya kodu"),
    ("kampanya_adi", "Kampanya adı"),
    ("magaza_kodu", "Mağaza kodu"),
    ("magaza_adi", "Mağaza adı"),
    ("aciklama", "Açıklama"),
    ("etiketler", "Etiketler"),
    ("bitis_sonrasi_url", "Bitiş sonrası URL"),
]
ALAN_ADLARI = dict(ALANLAR)

_TR = str.maketrans({"ç": "c", "ğ": "g", "ı": "i", "ö": "o", "ş": "s", "ü": "u", "â": "a", "î": "i", "û": "u"})


def normallestir(metin):
    """Büyük/küçük harf, Türkçe karakter, boşluk ve noktalama duyarsız karşılaştırma anahtarı."""
    metin = str(metin or "").replace("İ", "i").replace("I", "ı").lower().translate(_TR)
    return "".join(c for c in metin if c.isalnum())


_ESANLAMLI = {
    "kod": ["kod", "kisakod", "qrkod", "qrkodu", "kisakodu"],
    "ad": ["ad", "adi", "isim", "qradi", "qrkoduadi", "qrkodadi", "qrad"],
    "baslik": ["baslik", "hedefbasligi", "hedefbaslik", "title"],
    "tur": ["tur", "turu", "qrturu"],
    "hedef_url": ["hedefurl", "hedef", "url", "hedefadres", "hedefadresi", "hedefbaglanti", "yonlendirmeurl", "yonlendirmeadresi"],
    "sayfa": ["sayfa", "sayfaslug", "slug", "hedefsayfa", "sayfaadresi"],
    "baslangic": ["baslangic", "baslangictarihi", "baslama", "baslamatarihi", "baslangiczamani"],
    "bitis": ["bitis", "bitistarihi", "bitiszamani", "son", "sontarih"],
    "durum": ["durum"],
    "urun_kodu": ["urunkodu", "urunkod"],
    "urun_adi": ["urunadi", "urunad", "urun"],
    "kampanya_kodu": ["kampanyakodu", "kampanyakod"],
    "kampanya_adi": ["kampanyaadi", "kampanyaad", "kampanya"],
    "magaza_kodu": ["magazakodu", "magazakod", "subekodu"],
    "magaza_adi": ["magazaadi", "magazaad", "magaza", "subeadi", "sube"],
    "aciklama": ["aciklama", "not", "notlar"],
    "etiketler": ["etiketler", "etiket"],
    "bitis_sonrasi_url": ["bitissonrasiurl", "bitissonrasiadres", "bitissonrasiadresi", "bitisurl"],
}
_ESLEME_SOZLUGU = {}
for _alan, _adlar in _ESANLAMLI.items():
    for _ad in _adlar:
        _ESLEME_SOZLUGU[normallestir(_ad)] = _alan
for _alan, _etiket in ALANLAR:
    _ESLEME_SOZLUGU.setdefault(normallestir(_alan), _alan)
    _ESLEME_SOZLUGU.setdefault(normallestir(_etiket), _alan)


def otomatik_esleme(basliklar):
    """Her başlık için önerilen sistem alanı ("" = yok say). Aynı alan iki kez önerilmez."""
    kullanilan = set()
    sonuc = []
    for baslik in basliklar:
        alan = _ESLEME_SOZLUGU.get(normallestir(baslik), "")
        if alan in kullanilan:
            alan = ""
        if alan:
            kullanilan.add(alan)
        sonuc.append(alan)
    return sonuc


# --- Dosya okuma ---

class DosyaHatasi(Exception):
    """Kullanıcıya gösterilebilir dosya okuma hatası."""


def _hucre_metni(deger):
    if deger is None:
        return ""
    if isinstance(deger, datetime):
        if deger.time() == time(0, 0):
            return deger.strftime("%Y-%m-%d")
        return deger.strftime("%Y-%m-%d %H:%M")
    if isinstance(deger, date):
        return deger.strftime("%Y-%m-%d")
    if isinstance(deger, bool):
        return "true" if deger else "false"
    if isinstance(deger, float) and deger.is_integer():
        return str(int(deger))
    return str(deger).strip()


def _csv_coz(veri):
    try:
        metin = veri.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            metin = veri.decode("cp1254")
        except UnicodeDecodeError as hata:
            raise DosyaHatasi("CSV dosyasının karakter kodlaması okunamadı (UTF-8 veya Windows-1254 olmalı).") from hata
    ilk = next((s for s in metin.splitlines() if s.strip()), "")
    ayrac = max([";", ",", "\t"], key=lambda a: (ilk.count(a), a == ";"))
    okuyucu = csv.reader(io.StringIO(metin, newline=""), delimiter=ayrac)
    try:
        for satir in okuyucu:
            yield [h.strip() for h in satir]
    except csv.Error as hata:
        raise DosyaHatasi(f"CSV dosyası okunamadı: {hata}") from hata


_XLSX_RELS_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_XLSX_PKG_RELS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_XLSX_ACILAMADI = "Excel dosyası açılamadı; geçerli bir .xlsx dosyası olmalı."


def _ilk_sayfa_yolu(arsiv):
    """İlk sayfanın XML yolunu workbook.xml + rels üzerinden bulur; bulunamazsa ilk xl/worksheets/*.xml."""
    adlar = set(arsiv.namelist())
    try:
        if "xl/workbook.xml" in adlar and "xl/_rels/workbook.xml.rels" in adlar:
            kok = _guvenli_xml_kok(arsiv, "xl/workbook.xml")
            sayfa = next((e for e in kok.iter() if e.tag.rsplit("}", 1)[-1] == "sheet"), None)
            rid = sayfa.get(f"{{{_XLSX_RELS_NS}}}id") if sayfa is not None else None
            if rid:
                for r in _guvenli_xml_kok(arsiv, "xl/_rels/workbook.xml.rels"):
                    if r.get("Id") == rid and r.get("Target"):
                        hedef = r.get("Target")
                        yol = hedef.lstrip("/") if hedef.startswith("/") else posixpath.normpath(posixpath.join("xl", hedef))
                        if yol in adlar:
                            return yol
    except (ET.ParseError, DefusedXmlException, expat.ExpatError, DosyaHatasi):
        pass
    adaylar = sorted(a for a in adlar if a.startswith("xl/worksheets/") and a.endswith(".xml"))
    return adaylar[0] if adaylar else None


def _guvenli_xml_kok(arsiv, ad):
    veri = arsiv.read(ad)
    if len(veri) > 1024 * 1024:
        raise DosyaHatasi(_XLSX_ACILAMADI)
    if b"<!DOCTYPE" in veri or b"<!ENTITY" in veri:
        raise DosyaHatasi(_XLSX_ACILAMADI)
    try:
        return ET.fromstring(veri)
    except DefusedXmlException:
        raise DosyaHatasi(_XLSX_ACILAMADI) from None


def _sayfa_hucre_kontrol(arsiv, yol):
    """Sayfa XML'ini akış hâlinde tarar; toplam/satır başına hücre sayısını sınırlar.
    DOCTYPE/ENTITY içeren XML reddedilir (varlık genişletme saldırıları)."""
    en_cok_toplam = (EN_FAZLA_SATIR + EN_FAZLA_BOS_SATIR) * (EN_FAZLA_SUTUN + 1)
    en_cok_satir = EN_FAZLA_SUTUN * 4
    durum = {"toplam": 0, "satir": 0}

    class _Asiri(Exception):
        pass

    def basla(ad, _nitelik):
        yerel = ad.rsplit(" ", 1)[-1]
        if yerel == "row":
            durum["satir"] = 0
        elif yerel == "c":
            durum["toplam"] += 1
            durum["satir"] += 1
            if durum["toplam"] > en_cok_toplam or durum["satir"] > en_cok_satir:
                raise _Asiri

    def yasak(*_a):
        raise DosyaHatasi(_XLSX_ACILAMADI)

    ayristirici = expat.ParserCreate(namespace_separator=" ")
    ayristirici.StartElementHandler = basla
    ayristirici.StartDoctypeDeclHandler = yasak
    ayristirici.EntityDeclHandler = yasak
    try:
        with arsiv.open(yol) as akis:
            while True:
                parca = akis.read(65536)
                ayristirici.Parse(parca, not parca)
                if not parca:
                    break
    except _Asiri:
        raise DosyaHatasi(
            f"Excel sayfasında çok fazla hücre var (en fazla {EN_FAZLA_SATIR} satır, {EN_FAZLA_SUTUN} sütun; "
            "tekrarlanan veya boş hücreleri temizleyip yeniden kaydedin)."
        ) from None
    except (expat.ExpatError, zipfile.BadZipFile, OSError, ValueError, NotImplementedError, RuntimeError) as hata:
        raise DosyaHatasi(_XLSX_ACILAMADI) from hata


def _xlsx_boyut_kontrol(veri):
    """Zip arşivinin açılmış toplam boyutunu ve ilk sayfanın hücre sayısını OKUMADAN/openpyxl'e vermeden ÖNCE
    denetler (zip bombası / boş hücre şişirmesi / XML varlık genişletme)."""
    try:
        with zipfile.ZipFile(io.BytesIO(veri)) as arsiv:
            toplam = sum(b.file_size for b in arsiv.infolist())
            if toplam > EN_FAZLA_ACILMIS_BAYT:
                raise DosyaHatasi(f"Excel dosyasının açılmış boyutu en fazla {EN_FAZLA_ACILMIS_BAYT // (1024 * 1024)} MB olabilir.")
            yol = _ilk_sayfa_yolu(arsiv)
            if yol:
                _sayfa_hucre_kontrol(arsiv, yol)
    except (zipfile.BadZipFile, OSError, ValueError, NotImplementedError, RuntimeError) as hata:
        raise DosyaHatasi(_XLSX_ACILAMADI) from hata


def _xlsx_coz(veri):
    from openpyxl import load_workbook

    _xlsx_boyut_kontrol(veri)
    try:
        kitap = load_workbook(io.BytesIO(veri), read_only=True, data_only=True)
    except Exception as hata:
        raise DosyaHatasi("Excel dosyası açılamadı; geçerli bir .xlsx dosyası olmalı.") from hata
    try:
        if not kitap.worksheets:
            raise DosyaHatasi("Excel dosyasında sayfa yok.")
        okunan = 0
        # max_col: sütun sayısı sınırın bir fazlası; 60'tan fazla sütun hâlâ algılanır ama boş hücre şişirmesi olmaz.
        for satir in kitap.worksheets[0].iter_rows(max_col=EN_FAZLA_SUTUN + 1, values_only=True):
            okunan += 1
            if okunan > EN_FAZLA_SATIR + EN_FAZLA_BOS_SATIR:
                raise DosyaHatasi(f"Excel dosyası en fazla {EN_FAZLA_SATIR} veri satırı içerebilir (boş satırlar dahil toplam satır sınırı aşıldı).")
            yield [_hucre_metni(h) for h in satir]
    except DosyaHatasi:
        raise
    except Exception as hata:
        raise DosyaHatasi("Excel dosyası okunamadı.") from hata
    finally:
        kitap.close()


def dosya_oku(ad, veri):
    """(basliklar, satirlar) döndürür. satirlar: [{"no": dosyadaki satır no, "hucreler": [...]}].
    Tamamen boş satırlar atlanır. Sınır aşımında DosyaHatasi."""
    if len(veri) > EN_FAZLA_BAYT:
        raise DosyaHatasi(f"Dosya en fazla {EN_FAZLA_BAYT // (1024 * 1024)} MB olabilir.")
    kucuk = (ad or "").lower()
    if kucuk.endswith(".csv"):
        satirlar = _csv_coz(veri)
    elif kucuk.endswith(".xlsx"):
        satirlar = _xlsx_coz(veri)
    else:
        raise DosyaHatasi("Yalnızca .csv veya .xlsx dosyası yüklenebilir.")

    basliklar = None
    sonuc = []
    for no, hucreler in enumerate(satirlar, start=1):
        if not any(hucreler):
            continue
        if basliklar is None:
            while hucreler and not hucreler[-1]:
                hucreler.pop()
            if len(hucreler) > EN_FAZLA_SUTUN:
                raise DosyaHatasi(f"En fazla {EN_FAZLA_SUTUN} sütun desteklenir.")
            basliklar = hucreler
            continue
        if len(sonuc) >= EN_FAZLA_SATIR:
            raise DosyaHatasi(f"Dosya en fazla {EN_FAZLA_SATIR} veri satırı içerebilir.")
        hucreler = (hucreler + [""] * len(basliklar))[: len(basliklar)]
        sonuc.append({"no": no, "hucreler": hucreler})
    if not basliklar:
        raise DosyaHatasi("Dosya boş; ilk satır sütun başlıklarını içermelidir.")
    if not sonuc:
        raise DosyaHatasi("Dosyada başlık satırından sonra veri satırı yok.")
    return basliklar, sonuc


# --- Değer ayrıştırma ---

_TUR_ESLEME = {}
for _v, _e in QRKod.Tur.choices:
    _TUR_ESLEME[normallestir(_v)] = _v
    _TUR_ESLEME[normallestir(_e)] = _v
_TUR_ESLEME.update({normallestir("ürün"): "urun", normallestir("mağaza"): "magaza", normallestir("şube"): "magaza"})

_TARIH_BICIMLERI = ["%Y-%m-%d", "%d.%m.%Y"]
_TARIH_SAAT_BICIMLERI = ["%Y-%m-%d %H:%M", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d.%m.%Y %H:%M:%S", "%Y-%m-%dT%H:%M", "%Y-%m-%dT%H:%M:%S"]


def tarih_coz(metin, bitis=False):
    """Metni Europe/Istanbul saatinde aware datetime'a çevirir.
    Yalnızca tarih verilmişse: başlangıç gün başı 00:00; bitiş ertesi gün 00:00 (bitiş günü dahil)."""
    metin = metin.strip()
    for bicim in _TARIH_BICIMLERI:
        try:
            gun = datetime.strptime(metin, bicim).date()
        except ValueError:
            continue
        if bitis:
            gun += timedelta(days=1)
        return datetime.combine(gun, time.min, tzinfo=ZAMAN_DILIMI)
    for bicim in _TARIH_SAAT_BICIMLERI:
        try:
            return datetime.strptime(metin, bicim).replace(tzinfo=ZAMAN_DILIMI)
        except ValueError:
            continue
    raise ValueError("Tarih biçimi geçersiz (YYYY-AA-GG, GG.AA.YYYY veya YYYY-AA-GG SS:DD olmalı).")


def _hata_metni(hata):
    if hasattr(hata, "messages"):
        return " ".join(hata.messages)
    return str(hata)


# --- Satır doğrulama ---

@dataclass
class SatirSonucu:
    no: int
    kod: str = ""
    kod_otomatik: bool = False
    durum: str = "yeni"  # yeni | guncelle | degismedi | mevcut | hata
    hatalar: list = field(default_factory=list)  # [(alan, mesaj)]
    qr: object = None
    degisenler: list = field(default_factory=list)  # [(etiket, eski, yeni)]
    eski_hedef: str = ""
    yeni_hedef: str = ""
    iliskili: dict = field(default_factory=dict)  # {"urun": anahtar} yaratılacak ilişkili kayıtlar
    ad: str = ""
    # Mevcut kaydı güncellerken değişecek alanlar {nitelik: yeni değer}; uygulama sırasında kilitli taze satıra yazılır.
    guncelleme: dict = field(default_factory=dict)
    # Engellemeyen satır uyarıları (örn. aynı hedef başka kodlarda kullanılıyor).
    uyarilar: list = field(default_factory=list)

    def hata(self, alan, mesaj):
        self.hatalar.append((alan, mesaj))


@dataclass
class Analiz:
    satirlar: list
    yeni: list
    guncellenecek: list
    mevcut: list  # güncelleme adayı (değişiklik var)
    degismedi: list
    hatali: list
    iliskili_yeni: dict  # {("urun", "kod_kucuk"): nesne}
    iliskili_hata: dict = field(default_factory=dict)

    @property
    def ayni_hedef_sayisi(self):
        return sum(1 for s in self.satirlar if s.uyarilar)

    @property
    def hata_listesi(self):
        return [(s.no, alan, mesaj) for s in self.hatali for alan, mesaj in s.hatalar]


ILISKILI = {
    "urun": ("urun_kodu", "urun_adi", Urun, "Ürün"),
    "kampanya": ("kampanya_kodu", "kampanya_adi", Kampanya, "Kampanya"),
    "magaza": ("magaza_kodu", "magaza_adi", Magaza, "Mağaza"),
}


def _hedef_metni(hedef_turu, url, sayfa):
    if hedef_turu == QRKod.HedefTuru.SAYFA and sayfa is not None:
        return f"sayfa: {sayfa.slug}"
    return url or ""


def _yerel(dt):
    from django.utils import timezone

    return timezone.localtime(dt, ZAMAN_DILIMI).strftime("%d.%m.%Y %H:%M") if dt else ""


def analiz_et(ia, kalici=False):
    """İçe aktarma kaydındaki satırları eşleştirmeye göre doğrular. DB'ye yazmaz."""
    eslesme = ia.eslesme or {}
    alanlar = eslesme.get("alanlar", [])
    indeks = {a: i for i, a in enumerate(alanlar) if a}

    satirlar = ia.satirlar
    degerler = []
    for s in satirlar:
        h = s["hucreler"]
        degerler.append({a: (h[i].strip() if i < len(h) else "") for a, i in indeks.items()})

    # Dosya içi mükerrer kodlar (büyük/küçük harf duyarsız)
    kod_satirlari = {}
    for d, s in zip(degerler, satirlar):
        k = QRKod.anahtar(d.get("kod", ""))
        if k:
            kod_satirlari.setdefault(k, []).append(s["no"])

    # Veritabanı bağlamı
    anahtarlar = list(kod_satirlari)
    mevcut = {}
    for i in range(0, len(anahtarlar), 1000):
        for q in QRKod.objects.filter(kod_anahtari__in=anahtarlar[i:i + 1000]).select_related("sayfa", "urun", "kampanya", "magaza"):
            mevcut[q.kod_anahtari] = q
    sayfalar = {s.slug.lower(): s for s in KampanyaSayfasi.objects.all()}
    iliskili_mevcut = {tip: {o.kod.lower(): o for o in ayar[2].objects.all()} for tip, ayar in ILISKILI.items()}
    iliskili_yeni = {}
    iliskili_hata = {}
    kullanilan_kodlar = set(anahtarlar)
    ayar = GenelAyarlar.al()
    ayrilmis = ayar.ayrilmis_kume()
    # Otomatik kod: genel ayardaki biçime uyar. kalici=True yalnızca uygulama adımında (işlem içinde) sıralı sayacı
    # gerçekten ilerletir; önizlemede sayaç değişmez.
    kod_uret = None
    if ayar.kod_uretim_bicimi == GenelAyarlar.KodBicimi.SIRALI:  # rastgele biçim: aşağıdaki mevcut döngü aynen kullanılır
        kod_uret = kod_uretici(kalici=kalici, ayrilmis=ayrilmis, kullanilan=kullanilan_kodlar)
    # Aynı hedef uyarısı: dosyadaki URL'lerden hangileri arşivlenmemiş başka kodlarda kullanılıyor (tam eşleşme).
    ayni_hedef = {}
    if ayar.ayni_hedef_uyarisi:
        urller = sorted({d["hedef_url"] for d in degerler if d.get("hedef_url")})
        for i in range(0, len(urller), 1000):
            for url, kod in (
                QRKod.objects.filter(hedef_turu=QRKod.HedefTuru.HARICI, hedef_url__in=urller[i:i + 1000])
                .exclude(durum=QRKod.Durum.ARSIV).order_by("kod_anahtari").values_list("hedef_url", "kisa_kod")
            ):
                ayni_hedef.setdefault(url, []).append(kod)

    sonuclar = []
    for d, s in zip(degerler, satirlar):
        sonuclar.append(_satir_coz(
            s["no"], d, kod_satirlari, mevcut, sayfalar, iliskili_mevcut, iliskili_yeni, iliskili_hata, kullanilan_kodlar,
            ayrilmis=ayrilmis, ayni_hedef=ayni_hedef, kod_uret=kod_uret,
        ))

    analiz = Analiz(sonuclar, [], [], [], [], [], iliskili_yeni, iliskili_hata)
    for r in sonuclar:
        if r.durum == "hata":
            analiz.hatali.append(r)
        elif r.durum == "yeni":
            analiz.yeni.append(r)
        elif r.durum == "mevcut":
            analiz.mevcut.append(r)
        elif r.durum == "degismedi":
            analiz.degismedi.append(r)
    return analiz


def _satir_coz(no, d, kod_satirlari, mevcut, sayfalar, iliskili_mevcut, iliskili_yeni, iliskili_hata, kullanilan_kodlar,
              ayrilmis=frozenset(), ayni_hedef=None, kod_uret=None):
    r = SatirSonucu(no=no)
    kod = d.get("kod", "")
    anahtar = QRKod.anahtar(kod)
    r.kod = kod
    r.ad = d.get("ad", "")

    # Kod
    if kod:
        try:
            kisa_kod_dogrula(kod)
        except ValidationError as h:
            r.hata("kod", _hata_metni(h))
        if len(kod_satirlari.get(anahtar, [])) > 1:
            diger = ", ".join(str(n) for n in kod_satirlari[anahtar])
            r.hata("kod", f"Kod dosyada birden çok kez geçiyor (satırlar: {diger}).")
    eski = mevcut.get(anahtar) if kod else None
    if kod and eski is None:  # ayrılmış kelimeler yalnızca yeni kodlar için denetlenir; mevcut kayıtlar etkilenmez
        ayrilmis_hata = ayrilmis_kod_hatasi(kod, ayrilmis)
        if ayrilmis_hata:
            r.hata("kod", ayrilmis_hata)
    if eski is not None and eski.durum == QRKod.Durum.ARSIV:
        r.hata("kod", "Kod arşivde; arşivdeki kodlar içe aktarmayla güncellenemez.")

    # Tür
    tur = None
    if d.get("tur"):
        tur = _TUR_ESLEME.get(normallestir(d["tur"]))
        if tur is None:
            r.hata("tur", "Geçersiz tür. Geçerli değerler: " + ", ".join(f"{v} ({e})" for v, e in QRKod.Tur.choices) + ".")

    # Durum
    durum = None
    if d.get("durum"):
        n = normallestir(d["durum"])
        if n == "aktif":
            durum = QRKod.Durum.AKTIF
        elif n == "pasif":
            durum = QRKod.Durum.PASIF
        elif n in ("arsiv", "arsivli"):
            r.hata("durum", "Arşiv durumu içe aktarmayla verilemez (yalnızca aktif veya pasif).")
        else:
            r.hata("durum", "Geçersiz durum; aktif veya pasif olmalı.")

    # Hedef
    url = d.get("hedef_url", "")
    slug = d.get("sayfa", "")
    sayfa = None
    if url and slug:
        r.hata("hedef_url", "Hedef URL ve sayfa birlikte verilemez; yalnızca birini doldurun.")
    elif url:
        try:
            guvenli_url_dogrula(url)
        except ValidationError as h:
            r.hata("hedef_url", _hata_metni(h))
    elif slug:
        sayfa = sayfalar.get(slug.lower())
        if sayfa is None:
            r.hata("sayfa", f"Sayfa bulunamadı: {slug}")
        elif sayfa.durum == KampanyaSayfasi.Durum.ARSIV:
            r.hata("sayfa", f"Sayfa arşivde: {slug}")
    elif eski is None:
        r.hata("hedef_url", "Hedef zorunludur: hedef URL veya mevcut bir sayfa (slug) verin.")

    if d.get("bitis_sonrasi_url"):
        try:
            guvenli_url_dogrula(d["bitis_sonrasi_url"])
        except ValidationError as h:
            r.hata("bitis_sonrasi_url", _hata_metni(h))

    # Tarihler
    baslangic = bitis = None
    for alan in ("baslangic", "bitis"):
        if d.get(alan):
            try:
                deger = tarih_coz(d[alan], bitis=(alan == "bitis"))
            except ValueError as h:
                r.hata(alan, str(h))
                continue
            if alan == "baslangic":
                baslangic = deger
            else:
                bitis = deger
    if baslangic and bitis and bitis <= baslangic:
        r.hata("bitis", "Bitiş, başlangıçtan sonra olmalı.")

    # Ürün / kampanya / mağaza
    bagli = {}
    for tip, (kod_alani, ad_alani, model, etiket) in ILISKILI.items():
        ik, ia_ = d.get(kod_alani, ""), d.get(ad_alani, "")
        if not ik:
            if ia_:
                r.hata(kod_alani, f"{etiket} adı verilmiş ama {etiket.lower()} kodu boş.")
            continue
        var = iliskili_mevcut[tip].get(ik.lower())
        if var is not None:
            bagli[tip] = var
            continue
        if not ia_:
            r.hata(kod_alani, f"{etiket} kodu sistemde yok: {ik}. Oluşturulması için {etiket.lower()} adı da verilmeli.")
            continue
        anahtar_i = (tip, ik.lower())
        if anahtar_i not in iliskili_yeni:
            nesne = model(kod=ik, ad=ia_)
            try:
                nesne.full_clean(validate_unique=False)
            except ValidationError as h:
                iliskili_hata[anahtar_i] = _hata_metni(h)
            iliskili_yeni[anahtar_i] = nesne
        if anahtar_i in iliskili_hata:
            r.hata(kod_alani, f"{etiket} oluşturulamaz: {iliskili_hata[anahtar_i]}")
        else:
            r.iliskili[tip] = anahtar_i

    if r.hatalar:
        r.durum = "hata"
        return r

    # QR nesnesini kur
    if eski is None:
        qr = QRKod(kisa_kod=kod)
        if not kod:
            r.kod_otomatik = True
            if kod_uret is not None:
                uretilen = kod_uret()
            else:
                while True:
                    uretilen = otomatik_kisa_kod()
                    if uretilen.lower() not in kullanilan_kodlar and uretilen.lower() not in ayrilmis:
                        break
                kullanilan_kodlar.add(uretilen.lower())
            qr.kisa_kod = r.kod = uretilen
        qr.ad = d.get("ad", "")
        qr.baslik = d.get("baslik", "")
        qr.tur = tur or QRKod.Tur.KAMPANYA
        qr.durum = durum or QRKod.Durum.AKTIF
        qr.aciklama = d.get("aciklama", "")
        qr.etiketler = d.get("etiketler", "")
        qr.bitis_sonrasi_url = d.get("bitis_sonrasi_url", "")
        qr.baslangic, qr.bitis = baslangic, bitis
        if sayfa is not None:
            qr.hedef_turu, qr.sayfa = QRKod.HedefTuru.SAYFA, sayfa
        else:
            qr.hedef_turu, qr.hedef_url = QRKod.HedefTuru.HARICI, url
        for tip, var in bagli.items():
            setattr(qr, tip, var)
        r.yeni_hedef = _hedef_metni(qr.hedef_turu, qr.hedef_url, sayfa)
        r.durum = "yeni"
        if url and ayni_hedef and ayni_hedef.get(url):
            r.uyarilar.append("Bu hedef zaten şu kodlarda kullanılıyor: " + ", ".join(ayni_hedef[url][:5]))
    else:
        qr = eski
        r.kod = eski.kisa_kod
        r.eski_hedef = _hedef_metni(eski.hedef_turu, eski.hedef_url, eski.sayfa)
        degisen = []

        def ata(nitelik, etiket, yeni, goster_eski=None, goster_yeni=None):
            onceki = getattr(qr, nitelik)
            if onceki != yeni:
                degisen.append((etiket, goster_eski if goster_eski is not None else str(onceki or ""),
                                goster_yeni if goster_yeni is not None else str(yeni or "")))
                setattr(qr, nitelik, yeni)
                r.guncelleme[nitelik] = yeni

        # Boş hücre = "değiştirme"
        if d.get("ad"):
            ata("ad", "Ad", d["ad"])
        if d.get("baslik"):
            ata("baslik", "Başlık", d["baslik"])
        if tur:
            ata("tur", "Tür", tur, qr.get_tur_display(), QRKod.Tur(tur).label)
        if durum:
            ata("durum", "Durum", durum)
        if d.get("aciklama"):
            ata("aciklama", "Açıklama", d["aciklama"])
        if d.get("etiketler"):
            ata("etiketler", "Etiketler", d["etiketler"])
        if d.get("bitis_sonrasi_url"):
            ata("bitis_sonrasi_url", "Bitiş sonrası URL", d["bitis_sonrasi_url"])
        if baslangic:
            ata("baslangic", "Başlangıç", baslangic, _yerel(qr.baslangic), _yerel(baslangic))
        if bitis:
            ata("bitis", "Bitiş", bitis, _yerel(qr.bitis), _yerel(bitis))
        if url:
            if (qr.hedef_turu, qr.hedef_url) != (QRKod.HedefTuru.HARICI, url):
                degisen.append(("Hedef", r.eski_hedef, url))
            qr.hedef_turu, qr.hedef_url, qr.sayfa = QRKod.HedefTuru.HARICI, url, None
            r.guncelleme.update(hedef_turu=QRKod.HedefTuru.HARICI, hedef_url=url, sayfa=None)
        elif sayfa is not None:
            if (qr.hedef_turu, qr.sayfa_id) != (QRKod.HedefTuru.SAYFA, sayfa.pk):
                degisen.append(("Hedef", r.eski_hedef, f"sayfa: {sayfa.slug}"))
            qr.hedef_turu, qr.sayfa, qr.hedef_url = QRKod.HedefTuru.SAYFA, sayfa, ""
            r.guncelleme.update(hedef_turu=QRKod.HedefTuru.SAYFA, sayfa=sayfa, hedef_url="")
        for tip, var in bagli.items():
            onceki = getattr(qr, tip)
            if onceki is None or onceki.pk != var.pk:
                degisen.append((ILISKILI[tip][3], str(onceki.kod) if onceki else "", var.kod))
                setattr(qr, tip, var)
                r.guncelleme[tip] = var
        for tip, anahtar_i in r.iliskili.items():
            onceki = getattr(qr, tip)
            degisen.append((ILISKILI[tip][3], str(onceki.kod) if onceki else "", f"{anahtar_i[1]} (yeni)"))
        r.degisenler = degisen
        r.yeni_hedef = _hedef_metni(qr.hedef_turu, qr.hedef_url, qr.sayfa)
        r.durum = "mevcut" if degisen else "degismedi"

    r.qr = qr
    # Hızlı model denetimi (kayıt sırasında tam full_clean yeniden yapılır).
    if r.durum in ("yeni", "mevcut"):
        qr.kod_anahtari = QRKod.anahtar(qr.kisa_kod)
        try:
            qr.full_clean(
                exclude=["urun", "kampanya", "magaza", "sayfa", "olusturan", "degistiren"],
                validate_unique=False,
            )
        except ValidationError as h:
            if hasattr(h, "error_dict"):
                for alan, hatalar in h.message_dict.items():
                    for m in hatalar:
                        r.hata(alan if alan != "__all__" else "genel", m)
            else:
                r.hata("genel", _hata_metni(h))
            r.durum = "hata"
    return r
