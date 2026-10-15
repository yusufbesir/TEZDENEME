"""Güvenlik yeniden denetimi (Y1-Y4) testleri."""
import io
import re
import time
import zipfile
from io import StringIO

from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import Client, override_settings
from openpyxl import Workbook

from qr import ice_aktarma as modul
from qr.middleware import ip_anahtari
from qr.models import DenemeSayaci, DenetimKaydi
from qr.sayac import GirisDenemesi, ad_normalle

from .yardimci import PAROLA, TemelTest, betik_dizini

VARYANTLAR = [
    "editor", "ｅditor", "eｄitor", "edｉtor", "ediｔor", "editｏr", "editoｒ",
    "ｅｄitor", "\U0001d41editor", "\U0001d452ditor", "\U0001d486ditor", "\U0001d4eeditor",
]


class Y1UnicodeSayac(TemelTest):
    def post(self, ad, parola, ip):
        return Client().post("/panel/giris", {"username": ad, "password": parola}, HTTP_X_FORWARDED_FOR=ip)

    def test_ad_normalle(self):
        self.assertEqual(ad_normalle(" ｅditor "), "editor")
        self.assertEqual(ad_normalle("\U0001d41editor"), "editor")

    @override_settings(MFA_ZORUNLU=False)
    def test_unicode_varyantlar_hesap_kilidini_atlatamaz(self):
        kodlar = [self.post(VARYANTLAR[i % len(VARYANTLAR)], "yanlis", f"198.51.100.{i // 3}").status_code for i in range(30)]
        self.assertEqual(kodlar[:10], [200] * 10)
        self.assertTrue(all(k == 429 for k in kodlar[10:]), kodlar)
        # Doğru parola da kilit süresince reddedilir (varyant ve normal ad).
        self.assertEqual(self.post("ｅditor", PAROLA, "198.51.100.200").status_code, 429)
        self.assertEqual(self.post("editor", PAROLA, "198.51.100.201").status_code, 429)
        self.assertEqual(DenemeSayaci.objects.filter(anahtar__startswith="hesap:").count(), 1)

    @override_settings(MFA_ZORUNLU=False)
    def test_denetim_kaydinda_gercek_kullanici_adi(self):
        self.post("ｅditor", "yanlis", "198.51.100.1")
        kayit = DenetimKaydi.objects.filter(eylem="giris_basarisiz").latest("pk")
        self.assertIn("Denenen kullanıcı adı: editor", kayit.ozet)
        self.assertNotIn("bilinmeyen", kayit.ozet)
        self.post("yok-kullanici", "yanlis", "198.51.100.1")
        self.assertIn("bilinmeyen:", DenetimKaydi.objects.filter(eylem="giris_basarisiz").latest("pk").ozet)

    def test_hesap_anahtari_kullanici_pk_ile(self):
        a = GirisDenemesi("giris-deneme", "ｅditor", "1.2.3.4")
        b = GirisDenemesi("giris-deneme", "editor", "5.6.7.8")
        self.assertEqual(a.k_hesap, b.k_hesap)
        self.assertEqual(a.k_ui, GirisDenemesi("giris-deneme", "EDITOR", "1.2.3.4").k_ui)


@override_settings(MFA_ZORUNLU=False)
class Y3KisaDevre(TemelTest):
    def post(self, ad, parola, ip):
        return Client().post("/panel/giris", {"username": ad, "password": parola}, HTTP_X_FORWARDED_FOR=ip)

    def test_tek_ip_hesap_sayacini_sisiremez(self):
        for _ in range(25):
            self.post("yonetici", "yanlis", "203.0.113.9")
        sayac = DenemeSayaci.objects.get(anahtar=GirisDenemesi("x", "yonetici", "").k_hesap)
        self.assertLessEqual(sayac.sayi, 5)  # kullanıcı+IP sınırı (5) sonrası artmaz
        self.assertEqual(self.post("yonetici", PAROLA, "192.0.2.77").status_code, 302)

    def test_sayac_sifirla_komutu(self):
        for _ in range(12):
            self.post("editor", "yanlis", f"203.0.113.{_}")
        self.assertEqual(self.post("editor", PAROLA, "192.0.2.5").status_code, 429)
        cikti = StringIO()
        call_command("sayac_sifirla", "ｅditor", "--ip", "203.0.113.1", stdout=cikti)
        self.assertEqual(self.post("editor", PAROLA, "192.0.2.5").status_code, 302)
        kayit = DenetimKaydi.objects.filter(eylem="sayac_sifirlandi").latest("pk")
        self.assertIn("editor", kayit.ozet)
        self.assertIn("203.0.113.1", kayit.ozet)
        self.assertEqual(kayit.nesne_id, str(self.editor.pk))

    def test_sayac_sifirla_bilinmeyen_kullanici_hata_vermez(self):
        call_command("sayac_sifirla", "yok", stdout=StringIO())
        self.assertTrue(DenetimKaydi.objects.filter(eylem="sayac_sifirlandi").exists())


class Y4Ipv6(TemelTest):
    def test_ip_anahtari(self):
        self.assertEqual(ip_anahtari("203.0.113.9"), "203.0.113.9")
        self.assertEqual(ip_anahtari("2001:db8:1:2:aaaa:bbbb:cccc:dddd"), "2001:db8:1:2::/64")
        self.assertEqual(ip_anahtari("2001:db8:1:2:1::1"), ip_anahtari("2001:db8:1:2:ffff::2"))
        self.assertNotEqual(ip_anahtari("2001:db8:1:3::1"), ip_anahtari("2001:db8:1:2::1"))
        self.assertEqual(ip_anahtari("::ffff:1.2.3.4"), "1.2.3.4")
        self.assertEqual(ip_anahtari("bozuk"), "bozuk")

    @override_settings(MFA_ZORUNLU=False, IP_DENEME_SINIRI=6)
    def test_ipv6_64_icinde_dondurme_ip_sayacina_takilir(self):
        kodlar = []
        for i in range(10):
            r = Client().post("/panel/giris", {"username": f"yok{i}", "password": "x"}, HTTP_X_FORWARDED_FOR=f"2001:db8:5:5::{i + 1:x}")
            kodlar.append(r.status_code)
        self.assertEqual(kodlar[:6], [200] * 6)
        self.assertTrue(all(k == 429 for k in kodlar[6:]), kodlar)

    @override_settings(GIRIS_HIZ_SINIRI=3)
    def test_hiz_siniri_ipv6_64(self):
        kodlar = [Client().post("/panel/giris", {"username": "a", "password": "b"}, HTTP_X_FORWARDED_FOR=f"2001:db8:9:9::{i + 1:x}").status_code for i in range(5)]
        self.assertEqual(kodlar[3:], [429, 429])

    def test_yedek_dongusu_temizlik_icerir(self):
        betik = (betik_dizini() / "yedek-dongu.sh").read_text()
        self.assertIn("DELETE FROM qr_denemesayaci WHERE bitis < now()", betik)


def bomba_xlsx(hucre_sayisi=0, satir_hucre=None, doctype=False):
    kitap = Workbook()
    kitap.active.append(["kod", "hedef_url"])
    tampon = io.BytesIO()
    kitap.save(tampon)
    kaynak = zipfile.ZipFile(io.BytesIO(tampon.getvalue()))
    cikti = io.BytesIO()
    with zipfile.ZipFile(cikti, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
        for oge in kaynak.infolist():
            veri = kaynak.read(oge.filename)
            if oge.filename == "xl/worksheets/sheet1.xml":
                metin = re.sub(r"<dimension[^>]*/>", "", veri.decode())
                bas, son = metin.split("<sheetData>")
                govde = '<row r="1"><c r="A1" t="inlineStr"><is><t>kod</t></is></c></row>'
                if hucre_sayisi:
                    govde += '<row r="2">' + '<c r="A2" t="inlineStr"><is><t>x</t></is></c>' * hucre_sayisi + "</row>"
                if satir_hucre:
                    for r in range(2, 2 + satir_hucre[0]):
                        govde += f'<row r="{r}">' + '<c r="A1"/>' * satir_hucre[1] + "</row>"
                if doctype:
                    bas = bas.replace("?>", '?><!DOCTYPE x [<!ENTITY a "aaaa"><!ENTITY b "&a;&a;&a;&a;">]>', 1)
                veri = (bas + "<sheetData>" + govde + "</sheetData>" + son.split("</sheetData>", 1)[1]).encode()
            z.writestr(oge, veri)
    return cikti.getvalue()


class Y2XlsxBomba(TemelTest):
    def test_sinir_10mb(self):
        self.assertEqual(modul.EN_FAZLA_ACILMIS_BAYT, 10 * 1024 * 1024)

    def test_tek_satirda_milyonlarca_hucre_hizla_reddedilir(self):
        veri = bomba_xlsx(hucre_sayisi=1_000_000)  # açılmış ~42 MB (bomba2.py)
        t = time.time()
        with self.assertRaises(modul.DosyaHatasi):
            modul.dosya_oku("a.xlsx", veri)
        self.assertLess(time.time() - t, 2)

    def test_satir_basina_hucre_siniri_hizla_reddedilir(self):
        # 10 MB altında kalan ama bir satırda 241 hücre: reddedilmeli
        veri = bomba_xlsx(hucre_sayisi=modul.EN_FAZLA_SUTUN * 4 + 1)
        with self.assertRaisesMessage(modul.DosyaHatasi, "çok fazla hücre"):
            modul.dosya_oku("a.xlsx", veri)

    def test_toplam_hucre_siniri(self):
        toplam = (modul.EN_FAZLA_SATIR + modul.EN_FAZLA_BOS_SATIR) * (modul.EN_FAZLA_SUTUN + 1)
        # Aşmayı küçük sınırla dene: sabitleri geçici düşür
        with self.settings():
            eski = (modul.EN_FAZLA_SATIR, modul.EN_FAZLA_BOS_SATIR, modul.EN_FAZLA_SUTUN)
            modul.EN_FAZLA_SATIR, modul.EN_FAZLA_BOS_SATIR, modul.EN_FAZLA_SUTUN = 5, 5, 3  # toplam 40, satır başı 12
            try:
                with self.assertRaises(modul.DosyaHatasi):
                    modul.dosya_oku("a.xlsx", bomba_xlsx(satir_hucre=(10, 10)))  # 100 hücre
                self.assertGreater(toplam, 40)
            finally:
                modul.EN_FAZLA_SATIR, modul.EN_FAZLA_BOS_SATIR, modul.EN_FAZLA_SUTUN = eski

    def test_doctype_reddedilir(self):
        with self.assertRaises(modul.DosyaHatasi):
            modul.dosya_oku("a.xlsx", bomba_xlsx(doctype=True))

    def test_normal_xlsx_gecer(self):
        kitap = Workbook()
        kitap.active.append(["kod", "ad", "hedef_url"])
        kitap.active.append(["ABC-1", "Ad", "https://example.com/a"])
        tampon = io.BytesIO()
        kitap.save(tampon)
        basliklar, satirlar = modul.dosya_oku("a.xlsx", tampon.getvalue())
        self.assertEqual(basliklar, ["kod", "ad", "hedef_url"])
        self.assertEqual(len(satirlar), 1)

    def test_sayfa_yolu_workbook_rels_ile_bulunur(self):
        kitap = Workbook()
        kitap.active.append(["kod"])
        tampon = io.BytesIO()
        kitap.save(tampon)
        with zipfile.ZipFile(io.BytesIO(tampon.getvalue())) as z:
            self.assertEqual(modul._ilk_sayfa_yolu(z), "xl/worksheets/sheet1.xml")
