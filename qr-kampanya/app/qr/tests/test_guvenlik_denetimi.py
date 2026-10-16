"""Güvenlik denetimi bulguları için testler: MFA anahtarı, kilit atlatma, giriş DoS, XLSX DoS,
oturum sonlandırma, eşzamanlılık, takma adlı ziyaretçi özeti."""
import hashlib
import io
import re
import threading
import time
import zipfile
from datetime import date
from unittest import mock

import pyotp
from django.contrib.auth import SESSION_KEY
from django.contrib.sessions.models import Session
from django.core.management import call_command
from django.db import connection
from django.test import Client, TransactionTestCase, override_settings
from django.utils import timezone
from openpyxl import Workbook

from qr import ice_aktarma as ia_modul
from qr.models import DenemeSayaci, DenetimKaydi, HedefGecmisi, MFACihazi, QRKod
from qr.sayac import dene

from .test_ice_aktarma import IceAktarmaTabani, csv_dosyasi
from .test_mfa import SIMDI, T0, giris_yap
from .yardimci import PAROLA, TemelTest


def yanlis_giris(istemci, ad="editor", ip="203.0.113.1", parola="yanlis"):
    return istemci.post("/panel/giris", {"username": ad, "password": parola}, HTTP_X_FORWARDED_FOR=ip)


@override_settings(MFA_ZORUNLU=True)
class MFAAnahtarTesti(TemelTest):
    def kurulum_sayfasi(self, istemci):
        with mock.patch(SIMDI, return_value=T0):
            giris_yap(istemci)
            return istemci.get("/panel/mfa/kurulum").context["gizli_anahtar"]

    def test_saldirgan_ve_kurban_farkli_anahtar_gorur(self):
        a, b = Client(), Client()
        anahtar_a = self.kurulum_sayfasi(a)
        anahtar_b = self.kurulum_sayfasi(b)
        self.assertNotEqual(anahtar_a, anahtar_b)
        self.assertFalse(MFACihazi.objects.filter(kullanici=self.editor).exists())

    def test_saldirganin_anahtari_kurbanin_kurulumundan_sonra_gecersiz(self):
        saldirgan, kurban = Client(), Client()
        anahtar_s = self.kurulum_sayfasi(saldirgan)
        anahtar_k = self.kurulum_sayfasi(kurban)
        with mock.patch(SIMDI, return_value=T0):
            kurban.post("/panel/mfa/kurulum", {"kod": pyotp.TOTP(anahtar_k).at(int(T0))})
        self.assertIn(SESSION_KEY, kurban.session)
        cihaz = MFACihazi.objects.get(kullanici=self.editor)
        self.assertTrue(cihaz.onayli)
        self.assertEqual(cihaz.gizli_anahtar, anahtar_k)
        # Saldırgan kendi anahtarıyla kurulumu onaylayamaz...
        with mock.patch(SIMDI, return_value=T0 + 100):
            yanit = saldirgan.post("/panel/mfa/kurulum", {"kod": pyotp.TOTP(anahtar_s).at(int(T0 + 100))})
        self.assertNotIn(SESSION_KEY, saldirgan.session)
        self.assertNotEqual(yanit.status_code, 500)
        self.assertEqual(MFACihazi.objects.get(kullanici=self.editor).gizli_anahtar, anahtar_k)
        # ... ve yeni bir girişte kendi anahtarıyla MFA'yı da geçemez.
        yeni = Client()
        with mock.patch(SIMDI, return_value=T0 + 200):
            giris_yap(yeni)
            yeni.post("/panel/mfa/dogrula", {"kod": pyotp.TOTP(anahtar_s).at(int(T0 + 200))})
        self.assertNotIn(SESSION_KEY, yeni.session)

    def test_yanlis_kodla_veritabanina_cihaz_yazilmaz(self):
        istemci = Client()
        self.kurulum_sayfasi(istemci)
        istemci.post("/panel/mfa/kurulum", {"kod": "000000"})
        self.assertFalse(MFACihazi.objects.filter(kullanici=self.editor).exists())

    def test_eski_onaysiz_kayit_dikkate_alinmaz(self):
        eski = MFACihazi.objects.create(kullanici=self.editor, gizli_anahtar=pyotp.random_base32(), onayli=False)
        istemci = Client()
        anahtar = self.kurulum_sayfasi(istemci)
        self.assertNotEqual(anahtar, eski.gizli_anahtar)
        # Eski (sızmış olabilecek) anahtarın kodu kurulumu onaylamaz.
        with mock.patch(SIMDI, return_value=T0):
            istemci.post("/panel/mfa/kurulum", {"kod": pyotp.TOTP(eski.gizli_anahtar).at(int(T0))})
        self.assertNotIn(SESSION_KEY, istemci.session)
        with mock.patch(SIMDI, return_value=T0 + 100):
            istemci.post("/panel/mfa/kurulum", {"kod": pyotp.TOTP(anahtar).at(int(T0 + 100))})
        self.assertIn(SESSION_KEY, istemci.session)
        cihaz = MFACihazi.objects.get(kullanici=self.editor)
        self.assertEqual((cihaz.gizli_anahtar, cihaz.onayli), (anahtar, True))

    @override_settings(MFA_ZORUNLU=False)
    def test_hesabim_kurulumu_da_veritabaninda_onaysiz_anahtar_tutmaz(self):
        self.giris(self.editor)
        anahtar = self.client.get("/panel/hesap/mfa/kur").context["gizli_anahtar"]
        self.assertFalse(MFACihazi.objects.filter(kullanici=self.editor).exists())
        self.assertEqual(self.client.get("/panel/hesap/mfa/kur").context["gizli_anahtar"], anahtar)
        with mock.patch(SIMDI, return_value=T0):
            self.client.post("/panel/hesap/mfa/kur", {"kod": pyotp.TOTP(anahtar).at(int(T0))})
        self.assertEqual(MFACihazi.objects.get(kullanici=self.editor).gizli_anahtar, anahtar)
        self.assertTrue(MFACihazi.objects.get(kullanici=self.editor).onayli)


@override_settings(GIRIS_DENEME_SINIRI=3, HESAP_DENEME_SINIRI=5, IP_DENEME_SINIRI=6, GIRIS_KILIT_SURESI_SN=600)
class KilitAtlatmaTesti(TemelTest):
    def test_ip_dondurerek_hesap_kilidi_atlatilamaz(self):
        durumlar = [yanlis_giris(self.client, ip=f"198.51.100.{i}").status_code for i in range(5)]
        self.assertEqual(durumlar, [200] * 5)
        # Hesap sınırı doldu: yeni IP'den DOĞRU parola bile reddedilir.
        yanit = yanlis_giris(self.client, ip="198.51.100.99", parola=PAROLA)
        self.assertEqual(yanit.status_code, 429)
        self.assertNotIn(SESSION_KEY, self.client.session)
        # Mesaj genel: var olmayan hesapla aynı kilit metni.
        for i in range(5):
            yanlis_giris(self.client, ad="hic-yok", ip=f"198.51.100.{100 + i}")
        yok = yanlis_giris(self.client, ad="hic-yok", ip="198.51.100.200")
        self.assertEqual(yok.status_code, 429)
        self.assertEqual(yok.context["hata"], yanit.context["hata"])

    def test_hesap_kilidi_gecicidir(self):
        for i in range(5):
            yanlis_giris(self.client, ip=f"198.51.100.{i}")
        self.assertEqual(yanlis_giris(self.client, ip="198.51.100.50", parola=PAROLA).status_code, 429)
        DenemeSayaci.objects.update(bitis=timezone.now() - timezone.timedelta(seconds=1))
        self.assertEqual(yanlis_giris(self.client, ip="198.51.100.50", parola=PAROLA).status_code, 302)

    def test_baska_hesap_etkilenmez(self):
        for i in range(5):
            yanlis_giris(self.client, ip=f"198.51.100.{i}")
        self.assertEqual(yanlis_giris(self.client, ad="yonetici", ip="198.51.100.70", parola=PAROLA).status_code, 302)

    def test_kullanici_adi_dondurerek_ip_kilidi_atlatilamaz(self):
        for i in range(6):
            self.assertEqual(yanlis_giris(self.client, ad=f"kisi{i}", ip="192.0.2.77").status_code, 200)
        # IP sınırı doldu; yeni bir kullanıcı adı ve DOĞRU parola bile aynı IP'den reddedilir.
        self.assertEqual(yanlis_giris(self.client, ad="izleyici", ip="192.0.2.77", parola=PAROLA).status_code, 429)
        # Başka IP etkilenmez.
        self.assertEqual(yanlis_giris(self.client, ad="izleyici", ip="192.0.2.78", parola=PAROLA).status_code, 302)

    def test_basarili_giris_ip_ve_hesap_sayaclarini_sisirmez(self):
        for _ in range(4):
            c = Client()
            self.assertEqual(yanlis_giris(c, ad="izleyici", ip="192.0.2.90", parola=PAROLA).status_code, 302)
        self.assertEqual(DenemeSayaci.objects.filter(sayi__gt=0).count(), 0)

    @override_settings(MFA_ZORUNLU=True, GIRIS_DENEME_SINIRI=50, HESAP_DENEME_SINIRI=4, IP_DENEME_SINIRI=50)
    def test_mfa_denemeleri_hesap_sayacina_katilir_ve_yeniden_giris_sifirlamaz(self):
        MFACihazi.objects.create(kullanici=self.editor, gizli_anahtar=pyotp.random_base32(), onayli=True)
        durumlar, girisler = [], []
        with mock.patch(SIMDI, return_value=T0):
            for i in range(5):
                c = Client()  # her denemede farklı IP ve yeni oturum (parola doğru)
                girisler.append(yanlis_giris(c, ip=f"198.51.100.{i}", parola=PAROLA).status_code)
                durumlar.append(c.post("/panel/mfa/dogrula", {"kod": "000000"}, HTTP_X_FORWARDED_FOR=f"198.51.100.{i}").status_code)
        # İlk üç MFA hatası 200; dördüncüde hesap sınırı doldu (429). Beşinci parola girişi artık hiç kabul edilmez.
        self.assertEqual(durumlar[:4], [200, 200, 200, 429])
        self.assertEqual(girisler, [302, 302, 302, 302, 429])
        # Kilitliyken doğru parola da doğru kod da kabul edilmez.
        c = Client()
        yanit = yanlis_giris(c, ip="198.51.100.150", parola=PAROLA)
        self.assertEqual(yanit.status_code, 429)
        self.assertNotIn(SESSION_KEY, c.session)

    def test_sayac_atomik_artar(self):
        for beklenen in range(1, 8):
            izinli, sayi = dene("deneme-anahtari", 5, 600)
            self.assertEqual(sayi, beklenen)
            self.assertEqual(izinli, beklenen <= 5)
        self.assertEqual(DenemeSayaci.objects.get(anahtar="deneme-anahtari").sayi, 7)

    def test_sure_dolan_sayac_sifirlanir(self):
        for _ in range(3):
            dene("x", 5, 600)
        DenemeSayaci.objects.filter(anahtar="x").update(bitis=timezone.now() - timezone.timedelta(seconds=5))
        self.assertEqual(dene("x", 5, 600), (True, 1))


class SayacEszamanlilikTesti(TransactionTestCase):
    def _fixture_teardown(self):
        # Veritabanında silme koruması tetikleyicileri olduğundan tam flush yapılmaz; yalnızca bu testin verisi silinir.
        DenemeSayaci.objects.all().delete()

    def test_paralel_artirma_kayip_vermez(self):
        sonuclar = []

        def is_parcacigi():
            try:
                for _ in range(5):
                    sonuclar.append(dene("paralel", 1000, 600)[1])
            finally:
                connection.close()

        parcaciklar = [threading.Thread(target=is_parcacigi) for _ in range(8)]
        [p.start() for p in parcaciklar]
        [p.join() for p in parcaciklar]
        self.assertEqual(DenemeSayaci.objects.get(anahtar="paralel").sayi, 40)
        self.assertEqual(sorted(sonuclar), list(range(1, 41)))


@override_settings(GIRIS_HIZ_SINIRI=3)
class GirisDoSTesti(TemelTest):
    def setUp(self):
        super().setUp()
        DenemeSayaci.objects.filter(anahtar__startswith="hiz:").delete()

    def test_hiz_asilinca_429_ve_parola_dogrulamasi_calismaz(self):
        ip = {"HTTP_X_FORWARDED_FOR": "203.0.113.9"}
        for _ in range(3):
            self.assertEqual(self.client.post("/panel/giris", {"username": "editor", "password": "x"}, **ip).status_code, 200)
        with mock.patch("django.contrib.auth.forms.authenticate") as doğrula:
            yanit = self.client.post("/panel/giris", {"username": "editor", "password": PAROLA}, **ip)
        self.assertEqual(yanit.status_code, 429)
        doğrula.assert_not_called()
        self.assertEqual(yanit["Retry-After"], "60")
        # GET etkilenmez, başka IP etkilenmez.
        self.assertEqual(self.client.get("/panel/giris", **ip).status_code, 200)
        self.assertEqual(self.client.post("/panel/giris", {"username": "editor", "password": "x"}, HTTP_X_FORWARDED_FOR="203.0.113.10").status_code, 200)

    def test_mfa_uclari_da_sinirlanir(self):
        ip = {"HTTP_X_FORWARDED_FOR": "203.0.113.20"}
        for yol in ("/panel/mfa/dogrula", "/panel/mfa/kurulum"):
            DenemeSayaci.objects.filter(anahtar__startswith="hiz:").delete()
            kodlar = [self.client.post(yol, {"kod": "000000"}, **ip).status_code for _ in range(5)]
            self.assertEqual(kodlar[3:], [429, 429], yol)

    def test_sinir_kapaliyken_uygulanmaz(self):
        with override_settings(GIRIS_HIZ_SINIRI=0):
            kodlar = {self.client.post("/panel/giris", {"username": f"k{i}", "password": "x"}).status_code for i in range(15)}
        self.assertNotIn(429, kodlar)


class DenetimKaydiSinirTesti(TemelTest):
    @override_settings(GIRIS_DENEME_SINIRI=100, HESAP_DENEME_SINIRI=100, IP_DENEME_SINIRI=100)
    def test_ayni_ip_ve_kullanici_icin_kayit_sayisi_sinirlidir(self):
        for _ in range(15):
            yanlis_giris(self.client, ip="203.0.113.5")
        self.assertEqual(DenetimKaydi.objects.filter(eylem="giris_basarisiz").count(), 3)

    def test_var_olmayan_kullanici_adi_acik_yazilmaz(self):
        gizli_parola = "Yanlislikla-Yazilan-Parola-77!"
        yanlis_giris(self.client, ad=gizli_parola)
        yanlis_giris(self.client, ad="editor")
        kayitlar = list(DenetimKaydi.objects.filter(eylem="giris_basarisiz").values_list("ozet", flat=True))
        self.assertEqual(len(kayitlar), 2)
        self.assertFalse(any(gizli_parola in k for k in kayitlar))
        beklenen = "bilinmeyen:" + hashlib.sha256(gizli_parola.encode()).hexdigest()[:12]
        self.assertTrue(any(beklenen in k for k in kayitlar), kayitlar)
        self.assertTrue(any("editor" in k and "bilinmeyen" not in k for k in kayitlar), kayitlar)

    @override_settings(GIRIS_DENEME_SINIRI=2)
    def test_kilitlenme_kaydi_da_maskelenir(self):
        for _ in range(2):
            yanlis_giris(self.client, ad="Gizli-Parola-1")
        kayit = DenetimKaydi.objects.get(eylem="giris_kilitlendi")
        self.assertNotIn("Gizli-Parola-1", kayit.ozet)


def _bomba(satir_sayisi):
    """Denetçinin bomba.py'sine benzer: her satırda yalnızca XFD sütununda boş bir hücre."""
    b = io.BytesIO()
    kitap = Workbook()
    kitap.active.append(["kod", "hedef_url"])
    kitap.save(b)
    z = zipfile.ZipFile(io.BytesIO(b.getvalue()))
    cikti = io.BytesIO()
    with zipfile.ZipFile(cikti, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as zo:
        for it in z.infolist():
            veri = z.read(it.filename)
            if it.filename == "xl/worksheets/sheet1.xml":
                metin = re.sub(r"<dimension[^>]*/>", "", veri.decode())
                bas, kalan = metin.split("<sheetData>")
                ilk = '<row r="1"><c r="A1" t="inlineStr"><is><t>kod</t></is></c><c r="B1" t="inlineStr"><is><t>hedef_url</t></is></c></row>'
                govde = "".join(f'<row r="{i}"><c r="XFD{i}"/></row>' for i in range(2, satir_sayisi))
                veri = (bas + "<sheetData>" + ilk + govde + "</sheetData>" + kalan.split("</sheetData>", 1)[1]).encode()
            zo.writestr(it, veri)
    return cikti.getvalue()


class XlsxDoSTesti(TemelTest):
    def test_bos_hucre_bombasi_hizla_reddedilir(self):
        veri = _bomba(400_000)  # açılmış boyut sınırın altında (~15 MB); satır sınırı devreye girer
        t = time.time()
        with self.assertRaises(ia_modul.DosyaHatasi):
            ia_modul.dosya_oku("a.xlsx", veri)
        self.assertLess(time.time() - t, 2.0)

    def test_acilmis_boyut_siniri_okumadan_once_uygulanir(self):
        veri = _bomba(300_000)  # ~11 MB açılmış
        self.assertLess(len(veri), ia_modul.EN_FAZLA_BAYT)
        t = time.time()
        with mock.patch.object(ia_modul, "EN_FAZLA_ACILMIS_BAYT", 1024 * 1024):
            with self.assertRaisesMessage(ia_modul.DosyaHatasi, "açılmış boyutu"):
                ia_modul.dosya_oku("a.xlsx", veri)
        self.assertLess(time.time() - t, 2.0)

    def test_gecersiz_zip_nazikce_reddedilir(self):
        with self.assertRaisesMessage(ia_modul.DosyaHatasi, "geçerli bir .xlsx"):
            ia_modul.dosya_oku("a.xlsx", b"bu bir zip degil")
        with self.assertRaises(ia_modul.DosyaHatasi):
            ia_modul.dosya_oku("a.xlsx", b"PK\x03\x04" + b"\0" * 50)

    def test_cok_sutunlu_dosya_hala_algilanir(self):
        kitap = Workbook()
        kitap.active.append([f"s{i}" for i in range(ia_modul.EN_FAZLA_SUTUN + 5)])
        kitap.active.append(["x"] * 3)
        tampon = io.BytesIO()
        kitap.save(tampon)
        with self.assertRaisesMessage(ia_modul.DosyaHatasi, "sütun"):
            ia_modul.dosya_oku("a.xlsx", tampon.getvalue())

    def test_normal_dosya_calisir(self):
        kitap = Workbook()
        kitap.active.append(["kod", "hedef_url"])
        kitap.active.append([None, None])
        kitap.active.append(["A1", "https://example.com/a"])
        tampon = io.BytesIO()
        kitap.save(tampon)
        basliklar, satirlar = ia_modul.dosya_oku("a.xlsx", tampon.getvalue())
        self.assertEqual(basliklar, ["kod", "hedef_url"])
        self.assertEqual([s["no"] for s in satirlar], [3])


@override_settings(MFA_ZORUNLU=False)
class MFASifirlaOturumTesti(TemelTest):
    def hedef_oturumu_ac(self):
        c = Client()
        c.force_login(self.editor)
        MFACihazi.objects.create(kullanici=self.editor, gizli_anahtar=pyotp.random_base32(), onayli=True)
        return c

    def oturum_var_mi(self, istemci):
        return istemci.get("/panel/").status_code == 200

    def test_yonetici_sifirlayinca_hedefin_oturumlari_biter(self):
        c1, c2 = self.hedef_oturumu_ac(), Client()
        c2.force_login(self.editor)
        self.assertTrue(self.oturum_var_mi(c1) and self.oturum_var_mi(c2))
        self.giris(self.yonetici)
        self.client.post(f"/panel/kullanicilar/{self.editor.pk}/mfa-sifirla")
        self.assertFalse(MFACihazi.objects.filter(kullanici=self.editor).exists())
        self.assertEqual(c1.get("/panel/").status_code, 302)
        self.assertEqual(c2.get("/panel/").status_code, 302)
        self.assertEqual(self.client.get("/panel/").status_code, 200)  # yöneticinin kendi oturumu sürer

    def test_komut_oturumlari_sonlandirir(self):
        c = self.hedef_oturumu_ac()
        diger = Client()
        diger.force_login(self.izleyici)
        call_command("mfa_sifirla", "editor", stdout=io.StringIO())
        self.assertEqual(c.get("/panel/").status_code, 302)
        self.assertEqual(diger.get("/panel/").status_code, 200)

    def test_kendi_sifirlayan_mevcut_oturumunu_korur(self):
        cihaz = MFACihazi.objects.create(kullanici=self.editor, gizli_anahtar=pyotp.random_base32(), onayli=True)
        mevcut, baska = Client(), Client()
        mevcut.force_login(self.editor)
        baska.force_login(self.editor)
        with mock.patch(SIMDI, return_value=T0):
            mevcut.post("/panel/hesap/mfa/sifirla", {"kod": pyotp.TOTP(cihaz.gizli_anahtar).at(int(T0))})
        self.assertFalse(MFACihazi.objects.filter(kullanici=self.editor).exists())
        self.assertEqual(mevcut.get("/panel/").status_code, 200)
        self.assertEqual(baska.get("/panel/").status_code, 302)
        self.assertEqual(Session.objects.count(), 1)


class EszamanliDegisiklikTesti(IceAktarmaTabani):
    def test_toplu_durum_degisikligi_eszamanli_hedef_degisikligini_ezmez(self):
        from qr.servisler import qr_durum_degistir

        qr = self.qr_yap("ESZ001", "https://example.com/eski")
        eski_nesne = QRKod.objects.get(pk=qr.pk)  # yönetici listesinde önceden okunmuş
        QRKod.objects.filter(pk=qr.pk).update(hedef_url="https://example.com/baska-kullanici")  # eşzamanlı değişiklik
        qr_durum_degistir([eski_nesne], QRKod.Durum.PASIF, None)
        taze = QRKod.objects.get(pk=qr.pk)
        self.assertEqual(taze.durum, QRKod.Durum.PASIF)
        self.assertEqual(taze.hedef_url, "https://example.com/baska-kullanici")
        self.assertTrue(DenetimKaydi.objects.filter(eylem="qr_durum_degistir", nesne_id=str(qr.pk)).exists())

    def test_toplu_durum_degisikligi_hedef_gecmisi_yazmaz_hedef_ayni_kalirsa(self):
        from qr.servisler import qr_durum_degistir

        qr = self.qr_yap("ESZ002")
        once = HedefGecmisi.objects.filter(qr=qr).count()
        self.assertEqual(qr_durum_degistir([qr], QRKod.Durum.PASIF, None), 1)
        self.assertEqual(HedefGecmisi.objects.filter(qr=qr).count(), once)
        self.assertEqual(qr_durum_degistir([qr], QRKod.Durum.PASIF, None), 0)

    def test_ice_aktarma_guncellemesi_eszamanli_hedef_degisikligini_ezmez(self):
        qr = self.qr_yap("ESZ003", "https://example.com/ilk")
        ia = self.yukle(csv_dosyasi("kod,ad\nESZ003,Yeni Ad\n"))
        gercek = ia_modul.analiz_et

        def araya_gir(kayit, **kw):
            analiz = gercek(kayit, **kw)  # önizleme/analiz eski hedefi okudu
            QRKod.objects.filter(pk=qr.pk).update(hedef_url="https://example.com/baska-kullanici")
            return analiz

        with mock.patch("qr.panel.ice_aktarma.ia_modul.analiz_et", side_effect=araya_gir):
            self.assertEqual(self.uygula(ia, guncelle=True).status_code, 302)
        taze = QRKod.objects.get(pk=qr.pk)
        self.assertEqual(taze.ad, "Yeni Ad")
        self.assertEqual(taze.hedef_url, "https://example.com/baska-kullanici")

    def test_ice_aktarma_hedef_degisimi_gecmise_yazilir(self):
        qr = self.qr_yap("ESZ004", "https://example.com/ilk")
        ia = self.yukle(csv_dosyasi("kod,hedef_url\nESZ004,https://example.com/yeni\n"))
        self.uygula(ia, guncelle=True)
        self.assertEqual(QRKod.objects.get(pk=qr.pk).hedef_url, "https://example.com/yeni")
        g = HedefGecmisi.objects.filter(qr=qr).order_by("-pk").first()
        self.assertEqual((g.eski_hedef_url, g.yeni_hedef_url), ("https://example.com/ilk", "https://example.com/yeni"))


class ZiyaretciOzetiTesti(TemelTest):
    def test_gun_yerel_tarihtir(self):
        from qr.yonlendirme import ziyaretci_ozeti

        self.assertEqual(ziyaretci_ozeti("1.2.3.4", "ua"), ziyaretci_ozeti("1.2.3.4", "ua", gun=timezone.localdate()))
        with mock.patch("qr.yonlendirme.timezone.localdate", return_value=date(2030, 1, 1)):
            self.assertEqual(ziyaretci_ozeti("1.2.3.4", "ua"), ziyaretci_ozeti("1.2.3.4", "ua", gun=date(2030, 1, 1)))
        self.assertNotEqual(ziyaretci_ozeti("1.2.3.4", "ua", gun=date(2030, 1, 1)), ziyaretci_ozeti("1.2.3.4", "ua", gun=date(2030, 1, 2)))
