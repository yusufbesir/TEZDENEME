"""Panelden QR yönetimi, görseller, toplu işlem ve dışa aktarma."""
import csv
import io
import xml.etree.ElementTree as ET
import zipfile

import segno
from django.urls import reverse
from openpyxl import load_workbook
from PIL import Image

from qr.models import DenetimKaydi, QRKod
from qr.servisler import qr_png, qr_svg

from .yardimci import TemelTest

PNG_IMZASI = b"\x89PNG\r\n\x1a\n"


class QROlusturmaTesti(TemelTest):
    def test_panelden_qr_olusturma(self):
        self.giris(self.editor)
        yanit = self.client.post("/panel/qr/yeni", {
            "kisa_kod": "YENI-01", "ad": "Vitrin afişi", "tur": "kampanya", "hedef_turu": "harici",
            "hedef_url": "https://example.com/kampanya", "durum": "aktif", "etiketler": "vitrin, yaz",
        })
        qr = QRKod.objects.get(kisa_kod="YENI-01")
        self.assertRedirects(yanit, f"/panel/qr/{qr.pk}/")
        self.assertEqual(qr.olusturan, self.editor)
        self.assertEqual(qr.hedef_gecmisi.count(), 1)
        self.assertEqual(qr.etiket_listesi, ["vitrin", "yaz"])

    def test_denetim_kaydi_olusur(self):
        self.giris(self.editor)
        self.client.post("/panel/qr/yeni", {"kisa_kod": "DEN-01", "tur": "kampanya", "hedef_turu": "harici", "hedef_url": "https://example.com/d", "durum": "aktif"})
        kayit = DenetimKaydi.objects.get(eylem="qr_olustur")
        self.assertEqual(kayit.kullanici, self.editor)
        self.assertEqual(kayit.kullanici_adi, "editor")
        self.assertEqual(kayit.nesne_turu, "qrkod")

    def test_kampanya_urun_magaza_kayitlari_denetlenir(self):
        self.giris(self.editor)
        self.client.post("/panel/kampanyalar/yeni", {"kod": "YAZ26", "ad": "Yaz 2026", "aktif": "on"})
        self.client.post("/panel/urunler/yeni", {"kod": "U1", "ad": "Ürün 1"})
        self.client.post("/panel/magazalar/yeni", {"kod": "M1", "ad": "Merkez", "aktif": "on"})
        eylemler = set(DenetimKaydi.objects.values_list("eylem", flat=True))
        self.assertTrue({"kampanya_olustur", "urun_olustur", "magaza_olustur"} <= eylemler)

    def test_kampanya_bitis_baslangictan_once_reddedilir(self):
        self.giris(self.editor)
        yanit = self.client.post("/panel/kampanyalar/yeni", {
            "kod": "KK", "ad": "K", "baslangic": "2030-02-01T10:00", "bitis": "2030-01-01T10:00", "aktif": "on",
        })
        self.assertEqual(yanit.status_code, 200)
        self.assertTrue(yanit.context["form"].errors)

    def test_liste_arama_ve_filtre(self):
        k = self.kampanya_yap("KARA")
        self.qr_yap("ARA001", ad="Vitrin", kampanya=k, etiketler="afis,vitrin")
        self.qr_yap("ARA002", ad="Kasa", durum="pasif")
        self.giris(self.izleyici)
        self.assertContains(self.client.get("/panel/qr/?q=vitrin"), "ARA001")
        self.assertNotContains(self.client.get("/panel/qr/?q=vitrin"), "ARA002")
        self.assertNotContains(self.client.get("/panel/qr/?durum=pasif"), "ARA001")
        self.assertContains(self.client.get(f"/panel/qr/?kampanya={k.pk}"), "ARA001")
        self.assertNotContains(self.client.get(f"/panel/qr/?kampanya={k.pk}"), "ARA002")

    def test_liste_sayfalama_50(self):
        for i in range(55):
            self.qr_yap(f"SYF{i:03d}", hedef=f"https://example.com/{i}")
        self.giris(self.izleyici)
        yanit = self.client.get("/panel/qr/")
        self.assertEqual(len(yanit.context["sayfa"]), 50)
        self.assertEqual(len(self.client.get("/panel/qr/?sayfa_no=2").context["sayfa"]), 5)

    def test_detay_hedefi_test_et_baglantisi_guvenli(self):
        qr = self.qr_yap("DET001", "https://example.com/test")
        self.giris(self.izleyici)
        yanit = self.client.get(f"/panel/qr/{qr.pk}/")
        self.assertContains(yanit, 'href="https://example.com/test" target="_blank" rel="noopener noreferrer"')
        self.assertContains(yanit, f'src="/panel/qr/{qr.pk}/png"')
        self.assertContains(yanit, qr.kisa_url)

    def test_hatali_kayit_raporu(self):
        sayfa = self.sayfa_yap("taslak-r", durum="taslak")
        self.qr_yap("MUK001", "https://example.com/ayni")
        self.qr_yap("MUK002", "https://example.com/ayni")
        self.qr_yap("TEK001", "https://example.com/tek")
        bozuk = self.qr_yap("BOZ001", "https://example.com/x")
        QRKod.objects.filter(pk=bozuk.pk).update(hedef_url="http://127.0.0.1/ic")
        s = QRKod(kisa_kod="SYF001", hedef_turu="sayfa", sayfa=sayfa)
        from qr.servisler import qr_kaydet
        qr_kaydet(s, None)
        from .yardimci import gelecek
        self.qr_yap("GEC001", "https://example.com/gec", bitis=gelecek(-2))
        self.giris(self.izleyici)
        yanit = self.client.get("/panel/rapor/hatali")
        self.assertEqual([k for _, ql in yanit.context["mukerrer"] for k in [q.kisa_kod for q in ql]], ["MUK001", "MUK002"])
        self.assertEqual([q.kisa_kod for q in yanit.context["gecersiz_hedef"]], ["BOZ001"])
        self.assertEqual([q.kisa_kod for q in yanit.context["yayinda_olmayan_sayfa"]], ["SYF001"])
        self.assertEqual([q.kisa_kod for q in yanit.context["suresi_gecmis"]], ["GEC001"])


class GorselUclariTesti(TemelTest):
    def setUp(self):
        super().setUp()
        self.qr = self.qr_yap("GORSEL1", "https://example.com/gorsel-hedefi")

    def test_png_imzasi_ve_icerik_kisa_url(self):
        self.giris(self.izleyici)
        yanit = self.client.get(f"/panel/qr/{self.qr.pk}/png")
        self.assertEqual(yanit.status_code, 200)
        self.assertEqual(yanit["Content-Type"], "image/png")
        self.assertTrue(yanit.content.startswith(PNG_IMZASI))
        self.assertEqual(yanit.content, qr_png(self.qr.kisa_url))
        self.assertNotEqual(yanit.content, qr_png(self.qr.hedef_url))
        # PNG'deki modül matrisi, kısa URL'den üretilen matrisle aynı olmalı.
        matris = segno.make(self.qr.kisa_url, error="m", micro=False).matrix
        resim = Image.open(io.BytesIO(yanit.content)).convert("L")
        olcek, kenar = 10, 4
        self.assertEqual(resim.size, ((len(matris) + 2 * kenar) * olcek,) * 2)
        for s, satir in enumerate(matris):
            for c, koyu in enumerate(satir):
                piksel = resim.getpixel(((c + kenar) * olcek + 5, (s + kenar) * olcek + 5))
                self.assertEqual(piksel < 128, bool(koyu))

    def test_svg_gecerli_xml_ve_kisa_url(self):
        self.giris(self.izleyici)
        yanit = self.client.get(f"/panel/qr/{self.qr.pk}/svg")
        self.assertEqual(yanit.status_code, 200)
        self.assertEqual(yanit["Content-Type"], "image/svg+xml")
        kok = ET.fromstring(yanit.content)
        self.assertTrue(kok.tag.endswith("svg"))
        self.assertEqual(yanit.content, qr_svg(self.qr.kisa_url))
        self.assertNotEqual(yanit.content, qr_svg(self.qr.hedef_url))

    def test_hedef_degisse_gorsel_ayni_kalir(self):
        once = qr_png(self.qr.kisa_url)
        self.giris(self.editor)
        self.client.post(f"/panel/qr/{self.qr.pk}/duzenle", {
            "tur": "kampanya", "hedef_turu": "harici", "hedef_url": "https://example.com/baska", "durum": "aktif",
        })
        self.assertEqual(self.client.get(f"/panel/qr/{self.qr.pk}/png").content, once)

    def test_indirme_basligi(self):
        self.giris(self.izleyici)
        self.assertTrue(self.client.get(f"/panel/qr/{self.qr.pk}/png?indir=1")["Content-Disposition"].startswith("attachment"))
        self.assertTrue(self.client.get(f"/panel/qr/{self.qr.pk}/svg?indir=1")["Content-Disposition"].startswith("attachment"))
        self.assertTrue(self.client.get(f"/panel/qr/{self.qr.pk}/png")["Content-Disposition"].startswith("inline"))

    def test_giris_gerekli(self):
        for tur in ("png", "svg"):
            yanit = self.client.get(f"/panel/qr/{self.qr.pk}/{tur}")
            self.assertEqual(yanit.status_code, 302)
            self.assertIn("/panel/giris", yanit["Location"])

    def test_rolsuz_kullanici_403(self):
        self.giris(self.rolsuz)
        self.assertEqual(self.client.get(f"/panel/qr/{self.qr.pk}/png").status_code, 403)


class TopluIslemVeDisaAktarmaTesti(TemelTest):
    def setUp(self):
        super().setUp()
        self.a = self.qr_yap("TOP-A", "https://example.com/a", ad="=HYPERLINK(\"http://kotu.example\")")
        self.b = self.qr_yap("TOP-B", "https://example.com/b", ad="Normal ad")

    def test_zip_icerigi(self):
        self.giris(self.izleyici)
        yanit = self.client.post("/panel/qr/toplu", {"islem": "zip", "secili": [self.a.pk, self.b.pk]})
        self.assertEqual(yanit.status_code, 200)
        self.assertEqual(yanit["Content-Type"], "application/zip")
        arsiv = zipfile.ZipFile(io.BytesIO(yanit.content))
        self.assertEqual(sorted(arsiv.namelist()), ["TOP-A.png", "TOP-A.svg", "TOP-B.png", "TOP-B.svg", "liste.csv"])
        self.assertTrue(arsiv.read("TOP-A.png").startswith(PNG_IMZASI))
        self.assertEqual(arsiv.read("TOP-B.svg"), qr_svg(self.b.kisa_url))
        satirlar = list(csv.reader(io.StringIO(arsiv.read("liste.csv").decode("utf-8-sig")), delimiter=";"))
        self.assertEqual(satirlar[0], ["kisa_kod", "ad", "kisa_url", "hedef"])
        self.assertEqual(satirlar[1][0], "TOP-A")
        self.assertEqual(satirlar[1][2], self.a.kisa_url)
        self.assertEqual(satirlar[1][3], "https://example.com/a")
        self.assertTrue(satirlar[1][1].startswith("'="), "formül enjeksiyonu kaçışı")

    def test_zip_secim_yoksa_hata(self):
        self.giris(self.izleyici)
        yanit = self.client.post("/panel/qr/toplu", {"islem": "zip"})
        self.assertEqual(yanit.status_code, 302)

    def test_csv_formul_enjeksiyonu_ve_bom(self):
        self.giris(self.izleyici)
        yanit = self.client.get("/panel/qr/disa-aktar?bicim=csv")
        self.assertTrue(yanit.content.startswith(b"\xef\xbb\xbf"))
        satirlar = list(csv.reader(io.StringIO(yanit.content.decode("utf-8-sig")), delimiter=";"))
        adlar = [s[1] for s in satirlar[1:]]
        self.assertIn("'=HYPERLINK(\"http://kotu.example\")", adlar)
        self.assertNotIn("=HYPERLINK(\"http://kotu.example\")", adlar)
        self.assertEqual(satirlar[0][0], "Kısa kod")

    def test_csv_hucre_kacisi_tum_onekler(self):
        from qr.panel.yardimci import hucre_guvenli

        for onek in "=+-@":
            self.assertEqual(hucre_guvenli(f"{onek}cmd"), f"'{onek}cmd")
        self.assertEqual(hucre_guvenli("normal"), "normal")
        self.assertEqual(hucre_guvenli(5), 5)

    def test_xlsx_dis_aktarma_formul_degildir(self):
        self.giris(self.izleyici)
        yanit = self.client.get("/panel/qr/disa-aktar?bicim=xlsx")
        kitap = load_workbook(io.BytesIO(yanit.content))
        sayfa = kitap.active
        hucreler = {sayfa.cell(row=r, column=2) for r in range(2, sayfa.max_row + 1)}
        degerler = {h.value for h in hucreler}
        self.assertIn("'=HYPERLINK(\"http://kotu.example\")", degerler)
        for h in hucreler:
            self.assertEqual(h.data_type, "s")

    def test_toplu_pasif_ve_aktif(self):
        self.giris(self.editor)
        self.client.post("/panel/qr/toplu", {"islem": "pasif", "secili": [self.a.pk, self.b.pk]})
        self.assertEqual(QRKod.objects.filter(durum="pasif").count(), 2)
        self.client.post("/panel/qr/toplu", {"islem": "aktif", "secili": [self.a.pk]})
        self.a.refresh_from_db()
        self.assertEqual(self.a.durum, "aktif")
        self.assertTrue(DenetimKaydi.objects.filter(eylem="qr_durum_degistir").exists())

    def test_arsivleme_onay_ister(self):
        self.giris(self.editor)
        yanit = self.client.post("/panel/qr/toplu", {"islem": "arsivle", "secili": [self.a.pk]})
        self.assertEqual(yanit.status_code, 200)
        self.assertTemplateUsed(yanit, "qr/panel/qr_arsiv_onay.html")
        self.a.refresh_from_db()
        self.assertEqual(self.a.durum, "aktif")
        yanit = self.client.post("/panel/qr/toplu", {"islem": "arsivle", "secili": [self.a.pk], "onay": "1"})
        self.assertEqual(yanit.status_code, 302)
        self.a.refresh_from_db()
        self.assertEqual(self.a.durum, "arsiv")
        self.assertEqual(self.client.get("/k/TOP-A").status_code, 410)
        self.assertTrue(QRKod.objects.filter(pk=self.a.pk).exists())

    def test_izleyici_durum_degistiremez(self):
        self.giris(self.izleyici)
        yanit = self.client.post("/panel/qr/toplu", {"islem": "pasif", "secili": [self.a.pk]})
        self.assertEqual(yanit.status_code, 403)
        self.a.refresh_from_db()
        self.assertEqual(self.a.durum, "aktif")

    def test_toplu_yalnizca_post(self):
        self.giris(self.editor)
        self.assertEqual(self.client.get("/panel/qr/toplu").status_code, 405)
