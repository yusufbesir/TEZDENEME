"""XML ayrıştırma güvenliği (defusedxml): XXE, billion laughs ve harici varlık içeren .xlsx içe aktarmada reddedilir."""
import io
import re
import zipfile
from pathlib import Path

import openpyxl.xml
from defusedxml import DefusedXmlException
from django.conf import settings
from django.test import SimpleTestCase
from openpyxl import Workbook, load_workbook

from qr import ice_aktarma as modul

GIZLI_DOSYA = "/etc/passwd"
XXE = f'<!DOCTYPE x [<!ENTITY xxe SYSTEM "file://{GIZLI_DOSYA}">]>'
LAUGHS = (
    '<!DOCTYPE lolz [<!ENTITY lol "lol">'
    '<!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">'
    '<!ENTITY lol3 "&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;">'
    '<!ENTITY lol4 "&lol3;&lol3;&lol3;&lol3;&lol3;&lol3;&lol3;&lol3;&lol3;&lol3;">]>'
)


def temel_xlsx():
    kitap = Workbook()
    kitap.active.append(["kod", "hedef_url"])
    kitap.active.append(["ABC-1", "https://ornek.com"])
    tampon = io.BytesIO()
    kitap.save(tampon)
    return tampon.getvalue()


PAYLASILAN = "xl/sharedStrings.xml"
PAYLASILAN_ICERIK = (
    b'<?xml version="1.0" encoding="UTF-8"?><sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
    b'count="1" uniqueCount="1"><si><t>deger</t></si></sst>'
)


def degistir(veri, ad, donustur):
    """xlsx içindeki bir parçayı `donustur(bayt) -> bayt` ile değiştirir; parça yoksa (sharedStrings) eklenir."""
    if ad == PAYLASILAN:  # openpyxl bu parçayı yalnızca [Content_Types].xml'de bildirilmişse okur
        veri = degistir(veri, "[Content_Types].xml", lambda b: b.replace(
            b"</Types>", b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>'))
    kaynak = zipfile.ZipFile(io.BytesIO(veri))
    cikti = io.BytesIO()
    with zipfile.ZipFile(cikti, "w", zipfile.ZIP_DEFLATED) as z:
        for oge in kaynak.infolist():
            icerik = kaynak.read(oge.filename)
            z.writestr(oge, donustur(icerik) if oge.filename == ad else icerik)
        if ad not in kaynak.namelist() and ad == PAYLASILAN:
            z.writestr(ad, donustur(PAYLASILAN_ICERIK))
    return cikti.getvalue()


def doctype_ekle(doctype):
    def f(bayt):
        metin = bayt.decode("utf-8")
        metin = re.sub(r"(<\?xml[^>]*\?>)", lambda m: m.group(1) + doctype, metin, count=1)
        return metin.replace("<t>", "<t>&xxe;", 1).encode("utf-8") if "xxe" in doctype else metin.encode("utf-8")

    return f


class XmlGuvenligiTesti(SimpleTestCase):
    def test_openpyxl_defusedxml_kullaniyor(self):
        self.assertTrue(openpyxl.xml.DEFUSEDXML)

    def test_uygulama_kodunda_guvensiz_xml_ayristirici_yok(self):
        kok = Path(settings.BASE_DIR) / "qr"
        yasak = re.compile(r"^\s*(import xml\.(etree|dom|sax)|from xml\.(etree|dom|sax)|from xml import|import lxml|from lxml)", re.M)
        for yol in kok.rglob("*.py"):
            if "tests" in yol.parts or "migrations" in yol.parts:
                continue
            self.assertIsNone(yasak.search(yol.read_text(encoding="utf-8")), f"{yol}: doğrudan xml ayrıştırıcısı; defusedxml kullanın")

    def test_saglam_xlsx_hala_okunur(self):
        basliklar, satirlar = modul.dosya_oku("a.xlsx", temel_xlsx())
        self.assertEqual(basliklar, ["kod", "hedef_url"])
        self.assertEqual(satirlar[0]["hucreler"], ["ABC-1", "https://ornek.com"])

    def _reddedilir(self, veri):
        with self.assertRaises(modul.DosyaHatasi) as baglam:
            modul.dosya_oku("a.xlsx", veri)
        self.assertNotIn("root:", str(baglam.exception))  # sunucu dosyası içeriği sızmaz
        return baglam.exception

    def test_xxe_paylasilan_dizelerde_reddedilir_dosya_okunmaz(self):
        veri = degistir(temel_xlsx(), "xl/sharedStrings.xml", doctype_ekle(XXE))
        self._reddedilir(veri)

    def test_billion_laughs_paylasilan_dizelerde_reddedilir(self):
        veri = degistir(temel_xlsx(), "xl/sharedStrings.xml", doctype_ekle(LAUGHS))
        self._reddedilir(veri)

    def test_xxe_calisma_sayfasinda_reddedilir(self):
        veri = degistir(temel_xlsx(), "xl/worksheets/sheet1.xml", doctype_ekle(XXE))
        self._reddedilir(veri)

    def test_utf16_ile_gizlenmis_doctype_bayt_denetimini_asinca_da_reddedilir(self):
        # workbook.xml UTF-16: b"<!DOCTYPE" bayt araması bunu göremez; defusedxml/openpyxl yine de reddeder.
        def utf16(bayt):
            metin = bayt.decode("utf-8")
            metin = re.sub(r"<\?xml[^>]*\?>", "", metin, count=1)
            return ('<?xml version="1.0" encoding="UTF-16"?>' + XXE + metin).encode("utf-16")

        veri = degistir(temel_xlsx(), "xl/workbook.xml", utf16)
        self._reddedilir(veri)

    def test_openpyxl_dogrudan_xxe_yuklemeyi_reddeder(self):
        # Ön denetim atlansa bile (doğrudan openpyxl) varlık genişletme yapılmaz.
        veri = degistir(temel_xlsx(), "xl/sharedStrings.xml", doctype_ekle(XXE))
        with self.assertRaises(Exception) as baglam:
            kitap = load_workbook(io.BytesIO(veri), read_only=True)
            list(kitap.worksheets[0].iter_rows(values_only=True))
        self.assertNotIn("root:", str(baglam.exception))
        neden, zincir = baglam.exception, []
        while neden is not None:
            zincir.append(neden)
            neden = neden.__cause__ or neden.__context__
        self.assertTrue(any(isinstance(h, DefusedXmlException) for h in zincir), [type(h).__name__ for h in zincir])
