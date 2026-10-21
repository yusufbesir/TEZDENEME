"""YOURLS'ten eksik özellikler (1. parça): hedef başlığı, ayrılmış kodlar, aynı hedef uyarısı,
noindex/robots.txt ve QR oluşturma hız sınırı."""
import csv
import io
from unittest import mock

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from openpyxl import load_workbook

from qr import ice_aktarma as ia_modul
from qr.models import GenelAyarlar, IceAktarma, QRKod, ayrilmis_kodlar_kumesi

from .yardimci import TemelTest

NOINDEX = "noindex, nofollow"


def yeni_veri(**ek):
    veri = {"tur": "kampanya", "hedef_turu": "harici", "hedef_url": "https://example.com/hedef", "durum": "aktif"}
    veri.update(ek)
    return veri


def csv_yukle(istemci, metin):
    return istemci.post("/panel/qr/ice-aktar/yukle", {"dosya": SimpleUploadedFile("v.csv", metin.encode(), content_type="text/csv")})


class BaslikTesti(TemelTest):
    def test_baslik_olusturma_detay_ve_liste(self):
        self.giris(self.editor)
        self.client.post("/panel/qr/yeni", yeni_veri(kisa_kod="BSL-01", baslik="Sonbahar kampanyası sayfası"))
        qr = QRKod.objects.get(kisa_kod="BSL-01")
        self.assertEqual(qr.baslik, "Sonbahar kampanyası sayfası")
        self.assertContains(self.client.get(f"/panel/qr/{qr.pk}/"), "Sonbahar kampanyası sayfası")
        self.assertContains(self.client.get("/panel/qr/"), "Sonbahar kampanyası sayfası")

    def test_baslik_300_karakteri_asamaz(self):
        self.giris(self.editor)
        yanit = self.client.post("/panel/qr/yeni", yeni_veri(kisa_kod="BSL-02", baslik="x" * 301))
        self.assertEqual(yanit.status_code, 200)
        self.assertIn("baslik", yanit.context["form"].errors)
        self.assertFalse(QRKod.objects.filter(kisa_kod="BSL-02").exists())

    def test_baslik_duzenlenebilir(self):
        qr = self.qr_yap("BSL-03")
        self.giris(self.editor)
        self.client.post(f"/panel/qr/{qr.pk}/duzenle", yeni_veri(baslik="Yeni başlık", kisa_kod="BSL-03"))
        qr.refresh_from_db()
        self.assertEqual(qr.baslik, "Yeni başlık")

    def test_arama_basliga_bakar_ve_olmayani_getirmez(self):
        self.qr_yap("ARA-B1", baslik="Efsane indirim")
        self.qr_yap("ARA-B2", baslik="Başka şey")
        self.giris(self.yonetici)
        yanit = self.client.get("/panel/qr/", {"q": "efsane"})
        self.assertContains(yanit, "ARA-B1")
        self.assertNotContains(yanit, "ARA-B2")

    def test_csv_disa_aktarma_baslik_ve_formul_korumasi(self):
        self.qr_yap("DISA-1", baslik="=HYPERLINK(\"http://x\")")
        self.giris(self.yonetici)
        yanit = self.client.get("/panel/qr/disa-aktar", {"bicim": "csv"})
        satirlar = list(csv.reader(io.StringIO(yanit.content.decode("utf-8-sig")), delimiter=";"))
        self.assertEqual(satirlar[0][2], "Başlık")
        self.assertEqual(satirlar[1][2], "'=HYPERLINK(\"http://x\")")

    def test_xlsx_disa_aktarma_baslik_sutunu(self):
        self.qr_yap("DISA-2", baslik="+SUM(1)")
        self.giris(self.yonetici)
        yanit = self.client.get("/panel/qr/disa-aktar", {"bicim": "xlsx"})
        sayfa = load_workbook(io.BytesIO(yanit.content)).active
        self.assertEqual(sayfa.cell(1, 3).value, "Başlık")
        self.assertEqual(sayfa.cell(2, 3).value, "'+SUM(1)")

    def test_toplu_ice_aktarmada_baslik_istege_baglidir(self):
        self.giris(self.editor)
        csv_yukle(self.client, "kod,baslik,hedef_url\nIMP-B1,Başlıklı,https://example.com/b1\n")
        ia = IceAktarma.objects.latest("olusturma")
        self.client.post(f"/panel/qr/ice-aktar/{ia.pk}/uygula")
        self.assertEqual(QRKod.objects.get(kisa_kod="IMP-B1").baslik, "Başlıklı")
        # başlık sütunu olmayan dosya da çalışır
        csv_yukle(self.client, "kod,hedef_url\nIMP-B2,https://example.com/b2\n")
        ia = IceAktarma.objects.latest("olusturma")
        self.client.post(f"/panel/qr/ice-aktar/{ia.pk}/uygula")
        self.assertEqual(QRKod.objects.get(kisa_kod="IMP-B2").baslik, "")


class AyrilmisKodTesti(TemelTest):
    def test_kume_varsayilan_ve_ayardan(self):
        kume = ayrilmis_kodlar_kumesi("Kampanya, Yeni\nGizli")
        self.assertTrue({"admin", "panel", "robots", "favicon", "kampanya", "yeni", "gizli"} <= kume)

    def test_varsayilan_kelime_buyuk_kucuk_harf_reddedilir(self):
        self.giris(self.editor)
        for kod in ("admin", "ADMIN", "Panel", "api", "ROBOTS", "saglik"):
            yanit = self.client.post("/panel/qr/yeni", yeni_veri(kisa_kod=kod))
            self.assertEqual(yanit.status_code, 200, kod)
            self.assertIn("ayrılmış bir kelime", str(yanit.context["form"].errors["kisa_kod"]), kod)
            self.assertFalse(QRKod.objects.filter(kod_anahtari=kod.lower()).exists(), kod)

    def test_ayrilmamis_kod_kabul_edilir(self):
        self.giris(self.editor)
        self.client.post("/panel/qr/yeni", yeni_veri(kisa_kod="adminx"))
        self.assertTrue(QRKod.objects.filter(kisa_kod="adminx").exists())

    def test_ayardan_eklenen_kelime_reddedilir(self):
        ayar = GenelAyarlar.al()
        ayar.ayrilmis_kodlar = "kampanya\nyeni-urun, Gizli"
        ayar.save()
        self.giris(self.editor)
        for kod in ("KAMPANYA", "yeni-urun", "gizli"):
            yanit = self.client.post("/panel/qr/yeni", yeni_veri(kisa_kod=kod))
            self.assertEqual(yanit.status_code, 200, kod)
            self.assertFalse(QRKod.objects.filter(kod_anahtari=kod.lower()).exists(), kod)

    def test_mevcut_kayit_etkilenmez(self):
        # Ayrılmış kelime sonradan eklenir; mevcut kod hâlâ çalışır ve düzenlenebilir.
        qr = self.qr_yap("eskikod")
        ayar = GenelAyarlar.al()
        ayar.ayrilmis_kodlar = "eskikod"
        ayar.save()
        self.giris(self.editor)
        yanit = self.client.post(f"/panel/qr/{qr.pk}/duzenle", yeni_veri(kisa_kod="eskikod", ad="Yeni ad"))
        self.assertEqual(yanit.status_code, 302)
        qr.refresh_from_db()
        self.assertEqual(qr.ad, "Yeni ad")
        self.assertEqual(self.client.get("/k/eskikod").status_code, 302)

    def test_otomatik_uretim_ayrilmis_kodu_atlar(self):
        ayar = GenelAyarlar.al()
        ayar.ayrilmis_kodlar = "ABCD2345"
        ayar.save()
        self.giris(self.editor)
        with mock.patch("qr.panel.formlar.otomatik_kisa_kod", side_effect=["ABCD2345", "abcd2345", "SONUC234"]):
            self.client.post("/panel/qr/yeni", yeni_veri(kisa_kod=""))
        self.assertTrue(QRKod.objects.filter(kisa_kod="SONUC234").exists())
        self.assertFalse(QRKod.objects.filter(kod_anahtari="abcd2345").exists())

    def test_ayarlar_ekranindan_degistirilebilir(self):
        self.giris(self.yonetici)
        yanit = self.client.get("/panel/ayarlar")
        self.assertContains(yanit, "ayrilmis_kodlar")
        veri = {f: getattr(GenelAyarlar.al(), f) for f in ("marka_adi", "pasif_mesaj", "baslamadi_mesaj", "bitti_mesaj", "bulunamadi_mesaj", "yaklasan_bitis_gun")}
        veri.update(baslamadi_url="", bitti_url="", ayrilmis_kodlar="foo, bar", ayni_hedef_uyarisi="on")
        self.client.post("/panel/ayarlar", veri)
        self.assertEqual(GenelAyarlar.al().ayrilmis_kodlar, "foo, bar")

    def test_toplu_ice_aktarmada_ayrilmis_kod_reddedilir_digerleri_gecer(self):
        ayar = GenelAyarlar.al()
        ayar.ayrilmis_kodlar = "ozelkod"
        ayar.save()
        self.giris(self.editor)
        csv_yukle(self.client, "kod,hedef_url\nAdmin,https://example.com/1\nOZELKOD,https://example.com/2\nTAMAM-1,https://example.com/3\n")
        ia = IceAktarma.objects.latest("olusturma")
        analiz = ia_modul.analiz_et(ia)
        self.assertEqual(len(analiz.hatali), 2)
        self.assertTrue(all("ayrılmış" in h[0][1] for h in (s.hatalar for s in analiz.hatali)))
        self.assertEqual([s.kod for s in analiz.yeni], ["TAMAM-1"])

    def test_toplu_ice_aktarmada_otomatik_kod_ayrilmisi_atlar(self):
        ayar = GenelAyarlar.al()
        ayar.ayrilmis_kodlar = "ABCD2345"
        ayar.save()
        self.giris(self.editor)
        csv_yukle(self.client, "kod,hedef_url\n,https://example.com/1\n")
        ia = IceAktarma.objects.latest("olusturma")
        with mock.patch("qr.ice_aktarma.otomatik_kisa_kod", side_effect=["abcd2345", "SONUC234"]):
            analiz = ia_modul.analiz_et(ia)
        self.assertEqual(analiz.yeni[0].kod, "SONUC234")


class AyniHedefTesti(TemelTest):
    def test_ayni_hedef_ilk_gonderimde_uyarir_onayla_kaydeder(self):
        self.qr_yap("VAR-01", hedef="https://example.com/ortak")
        self.giris(self.editor)
        veri = yeni_veri(kisa_kod="YEN-01", hedef_url="https://example.com/ortak")
        yanit = self.client.post("/panel/qr/yeni", veri)
        self.assertEqual(yanit.status_code, 200)
        self.assertContains(yanit, "Bu hedef zaten şu kodlarda kullanılıyor: VAR-01")
        self.assertContains(yanit, "ayni_hedef_onay")
        self.assertFalse(QRKod.objects.filter(kisa_kod="YEN-01").exists())
        yanit = self.client.post("/panel/qr/yeni", {**veri, "ayni_hedef_onay": "on"})
        self.assertEqual(yanit.status_code, 302)
        self.assertTrue(QRKod.objects.filter(kisa_kod="YEN-01").exists())

    def test_farkli_hedefte_uyari_yok(self):
        self.qr_yap("VAR-02", hedef="https://example.com/ortak")
        self.giris(self.editor)
        yanit = self.client.post("/panel/qr/yeni", yeni_veri(kisa_kod="YEN-02", hedef_url="https://example.com/ortak/2"))
        self.assertEqual(yanit.status_code, 302)

    def test_arsivlenmis_kod_uyari_vermez(self):
        self.qr_yap("VAR-03", hedef="https://example.com/arsiv", durum="arsiv")
        self.giris(self.editor)
        yanit = self.client.post("/panel/qr/yeni", yeni_veri(kisa_kod="YEN-03", hedef_url="https://example.com/arsiv"))
        self.assertEqual(yanit.status_code, 302)

    def test_ayar_kapaliyken_uyari_yok(self):
        ayar = GenelAyarlar.al()
        ayar.ayni_hedef_uyarisi = False
        ayar.save()
        self.qr_yap("VAR-04", hedef="https://example.com/ortak")
        self.giris(self.editor)
        yanit = self.client.post("/panel/qr/yeni", yeni_veri(kisa_kod="YEN-04", hedef_url="https://example.com/ortak"))
        self.assertEqual(yanit.status_code, 302)

    def test_duzenlemede_uyari_yok(self):
        self.qr_yap("VAR-05", hedef="https://example.com/ortak")
        qr = self.qr_yap("VAR-06", hedef="https://example.com/baska")
        self.giris(self.editor)
        yanit = self.client.post(f"/panel/qr/{qr.pk}/duzenle", yeni_veri(kisa_kod="VAR-06", hedef_url="https://example.com/ortak"))
        self.assertEqual(yanit.status_code, 302)

    def test_toplu_ice_aktarmada_engel_yok_yalniz_uyari(self):
        self.qr_yap("VAR-07", hedef="https://example.com/ortak")
        self.giris(self.editor)
        csv_yukle(self.client, "kod,hedef_url\nIMP-A1,https://example.com/ortak\nIMP-A2,https://example.com/ayri\n")
        ia = IceAktarma.objects.latest("olusturma")
        analiz = ia_modul.analiz_et(ia)
        self.assertEqual(len(analiz.yeni), 2)
        self.assertEqual(analiz.ayni_hedef_sayisi, 1)
        self.assertIn("VAR-07", analiz.yeni[0].uyarilar[0])
        onizleme = self.client.get(f"/panel/qr/ice-aktar/{ia.pk}/onizle")
        self.assertContains(onizleme, "zaten başka kodlarda kullanılıyor")
        self.client.post(f"/panel/qr/ice-aktar/{ia.pk}/uygula")
        self.assertEqual(QRKod.objects.filter(hedef_url="https://example.com/ortak").count(), 2)

    def test_toplu_ice_aktarmada_ayar_kapaliyken_uyari_yok(self):
        ayar = GenelAyarlar.al()
        ayar.ayni_hedef_uyarisi = False
        ayar.save()
        self.qr_yap("VAR-08", hedef="https://example.com/ortak")
        self.giris(self.editor)
        csv_yukle(self.client, "kod,hedef_url\nIMP-C1,https://example.com/ortak\n")
        analiz = ia_modul.analiz_et(IceAktarma.objects.latest("olusturma"))
        self.assertEqual(analiz.ayni_hedef_sayisi, 0)


class NoindexVeRobotsTesti(TemelTest):
    def test_yonlendirme_302_noindex(self):
        self.qr_yap("NOI-01")
        yanit = self.client.get("/k/NOI-01")
        self.assertEqual(yanit.status_code, 302)
        self.assertEqual(yanit["X-Robots-Tag"], NOINDEX)

    def test_bilgi_sayfalari_200_404_410_noindex(self):
        self.qr_yap("NOI-02", durum="pasif")
        self.qr_yap("NOI-03", durum="arsiv")
        for yol, durum in (("/k/NOI-02", 200), ("/k/yok-boyle-kod", 404), ("/k/NOI-03", 410)):
            yanit = self.client.get(yol)
            self.assertEqual(yanit.status_code, durum, yol)
            self.assertEqual(yanit["X-Robots-Tag"], NOINDEX, yol)

    @override_settings(YONLENDIRME_HIZ_SINIRI=1)
    def test_429_yaniti_da_noindex(self):
        self.client.get("/k/yok-1")
        yanit = self.client.get("/k/yok-1")
        self.assertEqual(yanit.status_code, 429)
        self.assertEqual(yanit["X-Robots-Tag"], NOINDEX)

    def test_diger_halka_acik_uclarda_k_noindex_degil(self):
        yanit = self.client.get("/saglik")
        self.assertNotIn("X-Robots-Tag", yanit)

    def test_robots_txt(self):
        yanit = self.client.get("/robots.txt")
        self.assertEqual(yanit.status_code, 200)
        self.assertTrue(yanit["Content-Type"].startswith("text/plain"))
        self.assertEqual(yanit.content.decode(), "User-agent: *\nDisallow: /k/\nDisallow: /panel/\n")

    def test_robots_txt_yalniz_okuma(self):
        self.assertEqual(self.client.post("/robots.txt").status_code, 405)


class QROlusturmaHizSiniriTesti(TemelTest):
    @override_settings(QR_OLUSTURMA_HIZ_SINIRI=3)
    def test_sinir_asilinca_429_ve_turkce_mesaj(self):
        self.giris(self.editor)
        for i in range(3):
            yanit = self.client.post("/panel/qr/yeni", yeni_veri(kisa_kod=f"HIZ-0{i}", hedef_url=f"https://example.com/{i}"))
            self.assertEqual(yanit.status_code, 302, i)
        yanit = self.client.post("/panel/qr/yeni", yeni_veri(kisa_kod="HIZ-09", hedef_url="https://example.com/9"))
        self.assertEqual(yanit.status_code, 429)
        self.assertIn("Çok fazla QR kod oluşturma isteği", yanit.content.decode())
        self.assertEqual(yanit["Retry-After"], "60")
        self.assertFalse(QRKod.objects.filter(kisa_kod="HIZ-09").exists())

    @override_settings(QR_OLUSTURMA_HIZ_SINIRI=2)
    def test_sinir_kullanici_basinadir(self):
        self.giris(self.editor)
        for i in range(2):
            self.client.post("/panel/qr/yeni", yeni_veri(kisa_kod=f"HIZ-1{i}", hedef_url=f"https://example.com/a{i}"))
        self.assertEqual(self.client.post("/panel/qr/yeni", yeni_veri(kisa_kod="HIZ-19")).status_code, 429)
        self.giris(self.yonetici)
        yanit = self.client.post("/panel/qr/yeni", yeni_veri(kisa_kod="HIZ-20", hedef_url="https://example.com/yonetici"))
        self.assertEqual(yanit.status_code, 302)

    @override_settings(QR_OLUSTURMA_HIZ_SINIRI=2)
    def test_get_istekleri_sayilmaz(self):
        self.giris(self.editor)
        for _ in range(5):
            self.assertEqual(self.client.get("/panel/qr/yeni").status_code, 200)

    @override_settings(QR_OLUSTURMA_HIZ_SINIRI=2)
    def test_toplu_ice_aktarma_baslatma_ayni_sayaci_kullanir_satirlar_sayilmaz(self):
        self.giris(self.editor)
        # Çok satırlı tek dosya 1 sayılır (satır başına sayılmaz).
        satirlar = "".join(f"HIZ-S{i},https://example.com/s{i}\n" for i in range(10))
        yanit = csv_yukle(self.client, "kod,hedef_url\n" + satirlar)
        self.assertEqual(yanit.status_code, 302)
        ia = IceAktarma.objects.latest("olusturma")
        self.assertEqual(self.client.post(f"/panel/qr/ice-aktar/{ia.pk}/uygula").status_code, 302)
        self.assertEqual(QRKod.objects.filter(kisa_kod__startswith="HIZ-S").count(), 10)
        # Ikinci başlatma + tekil oluşturma aynı sayaçtan: toplam 2'yi aşınca 429.
        self.assertEqual(csv_yukle(self.client, "kod,hedef_url\nHIZ-T1,https://example.com/t1\n").status_code, 302)
        self.assertEqual(self.client.post("/panel/qr/yeni", yeni_veri(kisa_kod="HIZ-T2")).status_code, 429)
        self.assertEqual(csv_yukle(self.client, "kod,hedef_url\nHIZ-T3,https://example.com/t3\n").status_code, 429)

    @override_settings(QR_OLUSTURMA_HIZ_SINIRI=0)
    def test_sinir_sifirsa_kapali(self):
        self.giris(self.editor)
        for i in range(5):
            self.assertEqual(
                self.client.post("/panel/qr/yeni", yeni_veri(kisa_kod=f"HIZ-3{i}", hedef_url=f"https://example.com/k{i}")).status_code, 302
            )
