"""YOURLS'ten eksik özellikler (3b. parça): panel QR listesinde sıralama/filtre/sayfa boyutu, satır içi düzenleme
ve çöp kutusu (kalıcı silme YOK)."""
import csv
import io
from datetime import timedelta

from django.core.exceptions import ValidationError
from django.db import connection
from django.test import Client
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from qr.models import ApiAnahtari, DenetimKaydi, HedefGecmisi, QRKod, Tarama, Urun

from .yardimci import TEST_AYARLARI, TemelTest

LISTE = "/panel/qr/"


def kodlar(yanit):
    """Listede görünen kısa kodlar (sayfa sırasıyla)."""
    return [q.kisa_kod for q in yanit.context["sayfa"]]


class ListeTemel(TemelTest):
    def tara(self, qr, adet=1, bot=False):
        for _ in range(adet):
            Tarama.objects.create(qr=qr, sonuc="yonlendi", cihaz="bot" if bot else "mobil", bot=bot)

    def tarihle(self, qr, gun_once):
        QRKod.objects.filter(pk=qr.pk).update(olusturma=timezone.now() - timedelta(days=gun_once))

    def liste(self, **parametre):
        self.giris(self.editor)
        return self.client.get(LISTE, parametre)


class SiralamaTesti(ListeTemel):
    def setUp(self):
        super().setUp()
        self.a = self.qr_yap("CCC-01", "https://example.com/zz", ad="Beta", baslik="Kedi")
        self.b = self.qr_yap("AAA-01", "https://example.com/mm", ad="Gamma", baslik="Araba")
        self.c = self.qr_yap("BBB-01", "https://example.com/aa", ad="Alfa", baslik="Zebra")
        self.tarihle(self.a, 3)   # en eski
        self.tarihle(self.b, 2)
        self.tarihle(self.c, 1)   # en yeni
        self.tara(self.a, 1)
        self.tara(self.b, 3)
        self.tara(self.c, 2)

    def test_varsayilan_olusturma_azalan(self):
        self.assertEqual(kodlar(self.liste()), ["BBB-01", "AAA-01", "CCC-01"])

    def test_her_siralama_ve_yon(self):
        beklenen = {
            "kod": ["AAA-01", "BBB-01", "CCC-01"],
            "ad": ["BBB-01", "CCC-01", "AAA-01"],       # Alfa, Beta, Gamma
            "baslik": ["AAA-01", "CCC-01", "BBB-01"],   # Araba, Kedi, Zebra
            "hedef": ["BBB-01", "AAA-01", "CCC-01"],    # aa, mm, zz
            "olusturma": ["CCC-01", "AAA-01", "BBB-01"],
            "tarama": ["CCC-01", "BBB-01", "AAA-01"],   # 1, 2, 3
        }
        for alan, artan in beklenen.items():
            with self.subTest(alan=alan):
                self.assertEqual(kodlar(self.liste(sirala=alan, yon="artan")), artan)
                self.assertEqual(kodlar(self.liste(sirala=alan, yon="azalan")), artan[::-1])

    def test_gecersiz_deger_varsayilana_duser(self):
        varsayilan = kodlar(self.liste())
        self.assertEqual(kodlar(self.liste(sirala="yok", yon="yukari")), varsayilan)
        self.assertEqual(kodlar(self.liste(sirala="kod", yon="yukari")), ["CCC-01", "BBB-01", "AAA-01"])  # yon -> azalan

    def test_sql_enjeksiyon_denemesi_beyaz_listede_kalir(self):
        for kotu in ("id;drop", "id; DROP TABLE qr_qrkod;--", "kod_anahtari,(select 1)", "-olusturma", "tarama_sayisi desc", "?"):
            with self.subTest(kotu=kotu):
                with CaptureQueriesContext(connection) as sorgular:
                    yanit = self.liste(sirala=kotu, yon=kotu)
                self.assertEqual(yanit.status_code, 200)
                self.assertEqual(kodlar(yanit), ["BBB-01", "AAA-01", "CCC-01"])
                self.assertFalse(any("drop table" in s["sql"].lower() for s in sorgular.captured_queries))
        self.assertEqual(QRKod.objects.count(), 3)

    def test_baslik_baglantilari_siralar_ve_filtreleri_korur(self):
        yanit = self.liste(q="a", sirala="ad", yon="artan")
        basliklar = {b["etiket"]: b for b in yanit.context["basliklar"]}
        self.assertTrue(basliklar["Ad"]["mevcut"])
        self.assertIn("sirala=ad", basliklar["Ad"]["adres"])
        self.assertIn("yon=azalan", basliklar["Ad"]["adres"])  # aynı sütuna tekrar tıklama yönü çevirir
        self.assertIn("q=a", basliklar["Kısa kod"]["adres"])
        self.assertIn("sirala=kod", basliklar["Kısa kod"]["adres"])
        self.assertContains(yanit, 'href="?q=a&amp;sirala=kod&amp;yon=artan"')

    def test_tarama_sayisi_bot_haric_dogru(self):
        self.tara(self.a, 5, bot=True)
        sayilar = {q.kisa_kod: q.tarama_sayisi for q in self.liste().context["sayfa"]}
        self.assertEqual(sayilar, {"CCC-01": 1, "AAA-01": 3, "BBB-01": 2})
        bos = self.qr_yap("DDD-01", "https://example.com/bos")
        sayilar = {q.kisa_kod: q.tarama_sayisi for q in self.liste().context["sayfa"]}
        self.assertEqual(sayilar["DDD-01"], 0)
        self.assertEqual(bos.taramalar.count(), 0)

    def test_sorgu_sayisi_satir_sayisina_bagli_degil(self):
        self.giris(self.editor)
        with CaptureQueriesContext(connection) as az:
            self.client.get(LISTE)
        urun = Urun.objects.create(kod="U1", ad="Ürün 1")
        kamp = self.kampanya_yap()
        sayfa = self.sayfa_yap()
        for i in range(25):
            q = self.qr_yap(f"TOP-{i:02d}", f"https://example.com/t{i}", urun=urun, kampanya=kamp)
            self.tara(q, 2)
        # sayfa hedefli satır da N+1 yaratmamalı
        s = QRKod(kisa_kod="SYF-01", hedef_turu="sayfa", sayfa=sayfa)
        s.full_clean()
        from qr.servisler import qr_kaydet
        qr_kaydet(s, None)
        with CaptureQueriesContext(connection) as cok:
            yanit = self.client.get(LISTE, {"sirala": "tarama"})
        self.assertEqual(len(yanit.context["sayfa"]), 29)
        self.assertEqual(len(az), len(cok))


class FiltreTesti(ListeTemel):
    def setUp(self):
        super().setUp()
        self.urun = Urun.objects.create(kod="U1", ad="Ürün 1")
        self.diger = Urun.objects.create(kod="U2", ad="Ürün 2")
        self.a = self.qr_yap("ELMA-01", "https://example.com/armut", ad="Meyve", baslik="Sepet", urun=self.urun)
        self.b = self.qr_yap("ARMUT-1", "https://example.com/elma", ad="Elma kodu", baslik="Raf", urun=self.diger)
        self.c = self.qr_yap("KIRAZ-1", "https://example.com/kiraz", ad="Tatlı", baslik="Elma tezgâhı")
        self.tara(self.a, 1)
        self.tara(self.b, 4)
        self.tarihle(self.a, 10)
        self.tarihle(self.b, 5)
        self.tarihle(self.c, 1)

    def kume(self, **p):
        return sorted(kodlar(self.liste(**p)))

    def test_arama_alani(self):
        self.assertEqual(self.kume(q="elma"), ["ARMUT-1", "ELMA-01", "KIRAZ-1"])           # hepsi (kod, ad, başlık, hedef)
        self.assertEqual(self.kume(q="elma", alan="hepsi"), ["ARMUT-1", "ELMA-01", "KIRAZ-1"])
        self.assertEqual(self.kume(q="elma", alan="kod"), ["ELMA-01"])
        self.assertEqual(self.kume(q="elma", alan="ad"), ["ARMUT-1"])
        self.assertEqual(self.kume(q="elma", alan="baslik"), ["KIRAZ-1"])
        self.assertEqual(self.kume(q="elma", alan="hedef"), ["ARMUT-1"])
        self.assertEqual(self.kume(q="elma", alan="gecersiz"), ["ARMUT-1", "ELMA-01", "KIRAZ-1"])

    def test_tarama_sayisi_araligi(self):
        self.assertEqual(self.kume(tarama_min="1"), ["ARMUT-1", "ELMA-01"])
        self.assertEqual(self.kume(tarama_min="2"), ["ARMUT-1"])
        self.assertEqual(self.kume(tarama_max="1"), ["ELMA-01", "KIRAZ-1"])
        self.assertEqual(self.kume(tarama_max="0"), ["KIRAZ-1"])
        self.assertEqual(self.kume(tarama_min="1", tarama_max="3"), ["ELMA-01"])
        self.assertEqual(self.kume(tarama_min="abc", tarama_max="-1"), ["ARMUT-1", "ELMA-01", "KIRAZ-1"])  # geçersiz yok sayılır

    def test_bot_taramasi_sayiya_katilmaz(self):
        self.tara(self.c, 10, bot=True)
        self.assertEqual(self.kume(tarama_min="1"), ["ARMUT-1", "ELMA-01"])

    def test_olusturma_tarihi_once_sonra_arasi(self):
        bugun = timezone.localdate()
        gun = lambda n: (bugun - timedelta(days=n)).isoformat()
        self.assertEqual(self.kume(sonra=gun(6)), ["ARMUT-1", "KIRAZ-1"])
        self.assertEqual(self.kume(once=gun(6)), ["ELMA-01"])
        self.assertEqual(self.kume(sonra=gun(6), once=gun(3)), ["ARMUT-1"])
        self.assertEqual(self.kume(sonra=gun(5), once=gun(5)), ["ARMUT-1"])  # sınır günleri dahil
        self.assertEqual(self.kume(sonra="bozuk", once="2026-13-45"), ["ARMUT-1", "ELMA-01", "KIRAZ-1"])

    def test_urun_filtresi(self):
        self.assertEqual(self.kume(urun=str(self.urun.pk)), ["ELMA-01"])
        self.assertEqual(self.kume(urun=str(self.diger.pk)), ["ARMUT-1"])
        self.assertEqual(self.kume(urun="x"), ["ARMUT-1", "ELMA-01", "KIRAZ-1"])

    def test_sayfa_boyutu(self):
        for i in range(30):
            self.qr_yap(f"SYF-{i:02d}", f"https://example.com/s{i}")  # toplam 33
        self.assertEqual(self.liste().context["sayfa"].paginator.per_page, 50)
        for adet, beklenen in (("25", 25), ("50", 50), ("100", 100), ("200", 200), ("7", 50), ("abc", 50), ("-25", 50), ("", 50)):
            with self.subTest(adet=adet):
                yanit = self.liste(adet=adet)
                self.assertEqual(yanit.context["sayfa"].paginator.per_page, beklenen)
        self.assertEqual(len(self.liste(adet="25").context["sayfa"]), 25)
        self.assertEqual(len(self.liste(adet="25", sayfa_no="2").context["sayfa"]), 8)

    def test_filtreler_sayfalama_ve_disa_aktarmada_korunur(self):
        for i in range(30):
            q = self.qr_yap(f"ZZZ-{i:02d}", f"https://example.com/z{i}", ad="Toplu")
            self.tara(q, 2)
        parametre = dict(q="toplu", alan="ad", tarama_min="2", tarama_max="2", urun="", sirala="kod", yon="artan", adet="25")
        yanit = self.liste(**parametre)
        self.assertEqual(yanit.context["sayfa"].paginator.count, 30)
        icerik = yanit.content.decode()
        for parca in ("alan=ad", "tarama_min=2", "tarama_max=2", "sirala=kod", "yon=artan", "adet=25"):
            self.assertIn(parca, icerik)
            self.assertIn(parca, yanit.context["sorgu_dizesi"])
        self.assertIn("sayfa_no=2", icerik)
        self.assertRegex(icerik, r'href="/panel/qr/disa-aktar\?[^"]*tarama_min=2[^"]*bicim=csv"')
        self.assertRegex(icerik, r'href="/panel/qr/disa-aktar\?[^"]*sirala=kod[^"]*bicim=xlsx"')
        ikinci = self.client.get(LISTE, {**parametre, "sayfa_no": "2"})
        self.assertEqual(kodlar(ikinci), [f"ZZZ-{i:02d}" for i in range(25, 30)])
        # dışa aktarma aynı filtre ve sıralamayla
        csv_yanit = self.client.get("/panel/qr/disa-aktar", {**parametre, "bicim": "csv"})
        satirlar = list(csv.reader(io.StringIO(csv_yanit.content.decode("utf-8-sig")), delimiter=";"))
        self.assertEqual([s[0] for s in satirlar[1:]], [f"ZZZ-{i:02d}" for i in range(30)])
        xlsx = self.client.get("/panel/qr/disa-aktar", {**parametre, "bicim": "xlsx"})
        sayfa = load_workbook(io.BytesIO(xlsx.content)).active
        self.assertEqual(sayfa.max_row, 31)
        self.assertEqual(sayfa.cell(row=2, column=1).value, "ZZZ-00")

    def test_filtre_formu_secimleri_korur(self):
        yanit = self.liste(alan="ad", adet="100", sirala="tarama", yon="artan")
        self.assertContains(yanit, '<option value="ad" selected>')
        self.assertContains(yanit, '<option value="100" selected>')
        self.assertContains(yanit, '<input type="hidden" name="sirala" value="tarama">')
        self.assertContains(yanit, '<input type="hidden" name="yon" value="artan">')


class HizliDuzenleTesti(ListeTemel):
    def setUp(self):
        super().setUp()
        self.qr = self.qr_yap("DUZ-01", "https://example.com/eski", ad="Eski ad", baslik="Eski başlık")
        self.url = reverse("panel:qr_hizli_duzenle", args=[self.qr.pk])

    def gonder(self, veri, json=False, kullanici=None):
        self.giris(kullanici or self.editor)
        ek = {"HTTP_ACCEPT": "application/json"} if json else {}
        return self.client.post(self.url, veri, **ek)

    def test_yalnizca_post(self):
        self.giris(self.editor)
        self.assertEqual(self.client.get(self.url).status_code, 405)

    def test_yetki(self):
        for k in (self.izleyici, self.rolsuz):
            self.assertEqual(self.gonder({"ad": "x", "hedef_url": "https://example.com/x"}, kullanici=k).status_code, 403)
        anonim = Client().post(self.url, {"ad": "x"})
        self.assertEqual(anonim.status_code, 302)
        self.assertIn("/panel/giris", anonim["Location"])
        self.qr.refresh_from_db()
        self.assertEqual(self.qr.ad, "Eski ad")

    def test_csrf_zorunlu(self):
        istemci = Client(enforce_csrf_checks=True)
        istemci.force_login(self.editor)
        self.assertEqual(istemci.post(self.url, {"ad": "x", "hedef_url": "https://example.com/x"}).status_code, 403)
        self.qr.refresh_from_db()
        self.assertEqual(self.qr.ad, "Eski ad")

    def test_liste_csrf_tokenli_form_ve_js_olmadan_duzenleme_baglantisi(self):
        yanit = self.liste()
        icerik = yanit.content.decode()
        self.assertIn(f'action="{self.url}"', icerik)
        self.assertIn(f'href="/panel/qr/{self.qr.pk}/duzenle"', icerik)  # JS yoksa normal düzenleme sayfası
        self.assertIn(f'data-hizli-ac="hd-{self.qr.pk}"', icerik)
        self.assertIn("csrfmiddlewaretoken", icerik)
        self.assertNotIn("<script>", icerik)  # CSP: satır içi betik yok

    def test_basarili_json_ve_hedef_gecmisi_ve_denetim(self):
        yanit = self.gonder({"ad": "Yeni ad", "baslik": "Yeni başlık", "hedef_url": "https://example.com/yeni"}, json=True)
        self.assertEqual(yanit.status_code, 200)
        v = yanit.json()
        self.assertTrue(v["ok"])
        self.assertEqual((v["ad"], v["baslik"], v["hedef_url"]), ("Yeni ad", "Yeni başlık", "https://example.com/yeni"))
        self.qr.refresh_from_db()
        self.assertEqual((self.qr.ad, self.qr.baslik, self.qr.hedef_url), ("Yeni ad", "Yeni başlık", "https://example.com/yeni"))
        self.assertEqual(self.qr.degistiren, self.editor)
        gecmis = HedefGecmisi.objects.filter(qr=self.qr).order_by("-id").first()
        self.assertEqual((gecmis.eski_hedef_url, gecmis.yeni_hedef_url), ("https://example.com/eski", "https://example.com/yeni"))
        self.assertEqual(gecmis.kullanici, self.editor)
        self.assertEqual(HedefGecmisi.objects.filter(qr=self.qr).count(), 2)
        self.assertTrue(DenetimKaydi.objects.filter(eylem="qr_hedef_degistir", nesne_id=str(self.qr.pk), kullanici=self.editor).exists())

    def test_hedef_degismezse_gecmise_yazilmaz(self):
        yanit = self.gonder({"ad": "Sadece ad", "baslik": "Eski başlık", "hedef_url": "https://example.com/eski"}, json=True)
        self.assertTrue(yanit.json()["ok"])
        self.assertEqual(HedefGecmisi.objects.filter(qr=self.qr).count(), 1)
        self.assertFalse(DenetimKaydi.objects.filter(eylem="qr_hedef_degistir").exists())
        self.assertTrue(DenetimKaydi.objects.filter(eylem="qr_duzenle", nesne_id=str(self.qr.pk)).exists())

    def test_degisiklik_yoksa_hicbir_sey_yazilmaz(self):
        onceki = DenetimKaydi.objects.count()
        v = self.gonder({"ad": "Eski ad", "baslik": "Eski başlık", "hedef_url": "https://example.com/eski"}, json=True).json()
        self.assertTrue(v["ok"])
        self.assertFalse(v["degisti"])
        self.assertEqual(DenetimKaydi.objects.count(), onceki)

    def test_gecersiz_ve_ozel_hedefler_reddedilir(self):
        for kotu in (
            "javascript:alert(1)", "ftp://example.com/x", "not a url", "http://127.0.0.1/x", "http://192.168.1.5/x",
            "http://10.0.0.1/", "http://localhost/x", "https://user:parola@example.com/", "http://[::1]/",
        ):
            with self.subTest(kotu=kotu):
                yanit = self.gonder({"ad": "Y", "baslik": "", "hedef_url": kotu}, json=True)
                self.assertEqual(yanit.status_code, 400)
                v = yanit.json()
                self.assertFalse(v["ok"])
                self.assertIn("hedef_url", v["hatalar"])
        self.assertEqual(HedefGecmisi.objects.filter(qr=self.qr).count(), 1)
        self.qr.refresh_from_db()
        self.assertEqual((self.qr.ad, self.qr.hedef_url), ("Eski ad", "https://example.com/eski"))

    def test_bos_hedef_reddedilir(self):
        yanit = self.gonder({"ad": "Y", "hedef_url": ""}, json=True)
        self.assertEqual(yanit.status_code, 400)
        self.assertIn("hedef_url", yanit.json()["hatalar"])

    def test_yalniz_izinli_alanlar_degisir(self):
        self.gonder({
            "ad": "A", "baslik": "B", "hedef_url": "https://example.com/n", "durum": "pasif", "kisa_kod": "HACK-1",
            "tur": "magaza", "hedef_turu": "sayfa", "etiketler": "x",
        }, json=True)
        self.qr.refresh_from_db()
        self.assertEqual((self.qr.durum, self.qr.kisa_kod, self.qr.tur, self.qr.hedef_turu, self.qr.etiketler), ("aktif", "DUZ-01", "kampanya", "harici", ""))

    def test_sayfa_hedefli_qr_da_hedef_degismez(self):
        sayfa = self.sayfa_yap()
        qr = QRKod(kisa_kod="SYF-09", hedef_turu="sayfa", sayfa=sayfa, ad="Sayfa kodu")
        qr.full_clean()
        from qr.servisler import qr_kaydet
        qr_kaydet(qr, None)
        self.giris(self.editor)
        v = self.client.post(
            reverse("panel:qr_hizli_duzenle", args=[qr.pk]),
            {"ad": "Yeni sayfa kodu", "baslik": "B", "hedef_url": "https://example.com/saldiri"}, HTTP_ACCEPT="application/json",
        ).json()
        self.assertTrue(v["ok"])
        self.assertIsNone(v["hedef_url"])
        qr.refresh_from_db()
        self.assertEqual((qr.ad, qr.hedef_url, qr.sayfa_id, qr.hedef_turu), ("Yeni sayfa kodu", "", sayfa.pk, "sayfa"))
        self.assertEqual(HedefGecmisi.objects.filter(qr=qr).count(), 1)
        liste = self.client.get(LISTE).content.decode()
        satir = liste.split(f'id="hd-satir-{qr.pk}"')[1].split("</tr>")[0]
        self.assertNotIn('name="hedef_url"', satir)

    def test_arsiv_ve_cop_kaydi_reddedilir(self):
        for durum in ("arsiv", "cop"):
            with self.subTest(durum=durum):
                qr = self.qr_yap(f"RED-{durum}", "https://example.com/red", ad="Eski", durum=durum)
                self.giris(self.editor)
                yanit = self.client.post(
                    reverse("panel:qr_hizli_duzenle", args=[qr.pk]),
                    {"ad": "Yeni", "baslik": "", "hedef_url": "https://example.com/degisti"}, HTTP_ACCEPT="application/json",
                )
                self.assertEqual(yanit.status_code, 409)
                self.assertFalse(yanit.json()["ok"])
                qr.refresh_from_db()
                self.assertEqual((qr.ad, qr.hedef_url), ("Eski", "https://example.com/red"))
                self.assertEqual(HedefGecmisi.objects.filter(qr=qr).count(), 1)

    def test_normal_yanit_listeye_yonlendirir_ve_filtreyi_korur(self):
        yanit = self.gonder({
            "ad": "Normal yanıt", "baslik": "", "hedef_url": "https://example.com/n2",
            "donus_sorgu": "q=duz&sirala=ad&yon=artan&adet=25&sayfa_no=1",
        })
        self.assertEqual(yanit.status_code, 302)
        self.assertEqual(yanit["Location"], "/panel/qr/?q=duz&sirala=ad&yon=artan&adet=25&sayfa_no=1")
        self.assertIn("güncellendi", [str(m) for m in yanit.wsgi_request._messages][0])
        self.qr.refresh_from_db()
        self.assertEqual(self.qr.ad, "Normal yanıt")

    def test_normal_yanit_hata_mesaji_ve_yonlendirme(self):
        yanit = self.gonder({"ad": "x", "hedef_url": "http://127.0.0.1/x", "donus_sorgu": "q=duz"})
        self.assertEqual(yanit.status_code, 302)
        self.assertEqual(yanit["Location"], "/panel/qr/?q=duz")
        self.assertTrue(any("hedef_url" in str(m) for m in yanit.wsgi_request._messages))
        self.qr.refresh_from_db()
        self.assertEqual(self.qr.ad, "Eski ad")

    def test_donus_adresi_baska_siteye_gitmez(self):
        yanit = self.gonder({"ad": "x", "hedef_url": "https://example.com/e", "donus_sorgu": "//evil.example/?a=1"})
        self.assertTrue(yanit["Location"].startswith("/panel/qr/?"))
        self.assertNotIn("evil.example/", yanit["Location"].split("?")[0])

    def test_panel_js_hizli_duzenleme_ve_csp_uyumu(self):
        from pathlib import Path

        import qr as paket

        js = (Path(paket.__file__).parent / "static" / "qr" / "panel.js").read_text(encoding="utf-8")
        self.assertIn("data-hizli-form", js)
        self.assertIn("X-Requested-With", js)
        self.assertNotIn("eval(", js)
        self.assertNotIn("innerHTML", js)  # yanıt metinleri yalnızca textContent ile yazılır


class CopKutusuTesti(ListeTemel):
    def setUp(self):
        super().setUp()
        self.qr = self.qr_yap("COP-01", "https://example.com/cop", ad="Çöp adayı")
        self.pasif = self.qr_yap("COP-02", "https://example.com/cop2", ad="Pasif aday", durum="pasif")
        self.arsiv = self.qr_yap("COP-03", "https://example.com/cop3", ad="Arşiv aday", durum="arsiv")

    def cope_at(self, *qrler, kullanici=None):
        self.giris(kullanici or self.editor)
        veri = {"islem": "cope", "secili": [q.pk for q in qrler]}
        onay_sayfasi = self.client.post("/panel/qr/toplu", veri)
        self.assertEqual(onay_sayfasi.status_code, 200)
        return self.client.post("/panel/qr/toplu", {**veri, "onay": "1"})

    def test_onay_sayfasi_degisiklik_yapmaz(self):
        self.giris(self.editor)
        yanit = self.client.post("/panel/qr/toplu", {"islem": "cope", "secili": [self.qr.pk]})
        self.assertTemplateUsed(yanit, "qr/panel/qr_cop_onay.html")
        self.assertContains(yanit, "Kalıcı silme yoktur")
        self.qr.refresh_from_db()
        self.assertEqual(self.qr.durum, "aktif")

    def test_tasi_onceki_durum_ve_denetim(self):
        yanit = self.cope_at(self.qr, self.pasif, self.arsiv)
        self.assertEqual(yanit.status_code, 302)
        for q, onceki in ((self.qr, "aktif"), (self.pasif, "pasif"), (self.arsiv, "arsiv")):
            q.refresh_from_db()
            self.assertEqual((q.durum, q.onceki_durum), ("cop", onceki))
            self.assertIsNotNone(q.cope_atilma)
            self.assertEqual(q.degistiren, self.editor)
            kayit = DenetimKaydi.objects.get(eylem="qr_cope_tasi", nesne_id=str(q.pk))
            self.assertEqual(kayit.kullanici, self.editor)
            self.assertIn(onceki, kayit.ozet)
        self.assertEqual(QRKod.objects.count(), 3)  # silinmedi

    def test_detay_sayfasindan_tasima(self):
        self.giris(self.editor)
        detay = self.client.get(f"/panel/qr/{self.qr.pk}/")
        self.assertContains(detay, "Çöp kutusuna taşı")
        yanit = self.client.post("/panel/qr/toplu", {"islem": "cope", "secili": [self.qr.pk], "donus": str(self.qr.pk), "onay": "1"})
        self.assertRedirects(yanit, f"/panel/qr/{self.qr.pk}/")
        detay = self.client.get(f"/panel/qr/{self.qr.pk}/")
        self.assertContains(detay, "çöp kutusunda")
        self.assertContains(detay, "geri yükle")
        self.assertNotContains(detay, f"/panel/qr/{self.qr.pk}/duzenle")

    def test_yetki_ve_csrf(self):
        self.giris(self.izleyici)
        self.assertEqual(self.client.post("/panel/qr/toplu", {"islem": "cope", "secili": [self.qr.pk], "onay": "1"}).status_code, 403)
        istemci = Client(enforce_csrf_checks=True)
        istemci.force_login(self.editor)
        self.assertEqual(istemci.post("/panel/qr/toplu", {"islem": "cope", "secili": [self.qr.pk], "onay": "1"}).status_code, 403)
        self.assertEqual(self.client.get("/panel/qr/toplu").status_code, 405)
        self.qr.refresh_from_db()
        self.assertEqual(self.qr.durum, "aktif")

    def test_ana_listede_gorunmez_durum_filtresiyle_gorunur(self):
        self.cope_at(self.qr)
        self.assertEqual(sorted(kodlar(self.liste())), ["COP-02", "COP-03"])
        self.assertEqual(kodlar(self.liste(durum="cop")), ["COP-01"])
        disa = self.client.get("/panel/qr/disa-aktar", {"bicim": "csv"}).content.decode("utf-8-sig")
        self.assertNotIn("COP-01", disa)
        self.assertIn("COP-01", self.client.get("/panel/qr/disa-aktar", {"bicim": "csv", "durum": "cop"}).content.decode("utf-8-sig"))

    def test_cop_kutusu_listesi(self):
        self.cope_at(self.qr, self.pasif)
        self.giris(self.izleyici)
        yanit = self.client.get("/panel/cop-kutusu/")
        self.assertEqual(yanit.status_code, 200)
        self.assertEqual(sorted(q.kisa_kod for q in yanit.context["sayfa"]), ["COP-01", "COP-02"])
        self.assertNotContains(yanit, "geri-yukle")  # izleyici geri yükleyemez
        self.giris(self.editor)
        yanit = self.client.get("/panel/cop-kutusu/")
        self.assertContains(yanit, "Geri yükle")
        self.assertNotContains(yanit, "Kalıcı olarak sil")
        self.giris(self.rolsuz)
        self.assertEqual(self.client.get("/panel/cop-kutusu/").status_code, 403)

    def test_tekil_geri_yukleme_onceki_duruma_doner(self):
        self.cope_at(self.qr, self.pasif, self.arsiv)
        self.giris(self.editor)
        for q, beklenen in ((self.pasif, "pasif"), (self.arsiv, "arsiv"), (self.qr, "aktif")):
            yanit = self.client.post("/panel/cop-kutusu/geri-yukle", {"secili": [q.pk]})
            self.assertRedirects(yanit, "/panel/cop-kutusu/")
            q.refresh_from_db()
            self.assertEqual((q.durum, q.onceki_durum, q.cope_atilma), (beklenen, "", None))
            kayit = DenetimKaydi.objects.get(eylem="qr_copten_geri_yukle", nesne_id=str(q.pk))
            self.assertIn(beklenen, kayit.ozet)

    def test_toplu_geri_yukleme(self):
        self.cope_at(self.qr, self.pasif)
        self.client.post("/panel/cop-kutusu/geri-yukle", {"secili": [self.qr.pk, self.pasif.pk, self.arsiv.pk]})
        self.qr.refresh_from_db(), self.pasif.refresh_from_db(), self.arsiv.refresh_from_db()
        self.assertEqual((self.qr.durum, self.pasif.durum, self.arsiv.durum), ("aktif", "pasif", "arsiv"))
        self.assertEqual(DenetimKaydi.objects.filter(eylem="qr_copten_geri_yukle").count(), 2)  # çöpte olmayan etkilenmez

    def test_geri_yukleme_yetki_csrf_ve_sinir(self):
        self.cope_at(self.qr)
        self.giris(self.izleyici)
        self.assertEqual(self.client.post("/panel/cop-kutusu/geri-yukle", {"secili": [self.qr.pk]}).status_code, 403)
        istemci = Client(enforce_csrf_checks=True)
        istemci.force_login(self.editor)
        self.assertEqual(istemci.post("/panel/cop-kutusu/geri-yukle", {"secili": [self.qr.pk]}).status_code, 403)
        self.giris(self.editor)
        self.assertEqual(self.client.get("/panel/cop-kutusu/geri-yukle").status_code, 405)
        yanit = self.client.post("/panel/cop-kutusu/geri-yukle", {"secili": list(range(1, 1002))}, follow=True)
        self.assertContains(yanit, "en fazla 1000")
        self.qr.refresh_from_db()
        self.assertEqual(self.qr.durum, "cop")

    def test_toplu_sinir_1000(self):
        self.giris(self.editor)
        yanit = self.client.post("/panel/qr/toplu", {"islem": "cope", "secili": list(range(1, 1002)), "onay": "1"}, follow=True)
        self.assertContains(yanit, "en fazla 1000")

    def test_cop_kod_k_altinda_410_ve_tarama_sonucu(self):
        self.cope_at(self.qr)
        yanit = Client().get("/k/COP-01")
        self.assertEqual(yanit.status_code, 410)
        self.assertNotIn("Location", yanit)
        self.assertEqual(Tarama.objects.get(qr=self.qr).sonuc, "cop")
        arsiv = Client().get("/k/COP-03")
        self.assertEqual(arsiv.status_code, 410)
        # geri yüklenince tekrar yönlenir
        self.client.post("/panel/cop-kutusu/geri-yukle", {"secili": [self.qr.pk]})
        self.assertEqual(Client().get("/k/COP-01").status_code, 302)

    def test_cop_kodu_baskasina_verilemez(self):
        self.cope_at(self.qr)
        self.giris(self.editor)
        yanit = self.client.post("/panel/qr/yeni", {
            "kisa_kod": "cop-01", "tur": "kampanya", "hedef_turu": "harici", "hedef_url": "https://example.com/baska", "durum": "aktif",
        })
        self.assertEqual(yanit.status_code, 200)
        self.assertContains(yanit, "zaten kullanılıyor")
        self.assertEqual(QRKod.objects.filter(kod_anahtari="cop-01").count(), 1)

    def test_cop_hedef_tutarlilik_kisiti_muaf(self):
        from django.db import IntegrityError, transaction

        self.cope_at(self.qr)
        QRKod.objects.filter(pk=self.qr.pk).update(hedef_url="")  # çöpteki kayıtta boş hedef kısıta takılmaz
        with self.assertRaises(IntegrityError), transaction.atomic():
            QRKod.objects.filter(pk=self.pasif.pk).update(hedef_url="")  # pasif için kısıt sürer

    def test_cop_durumu_normal_formdan_verilemez_ve_cop_duzenlenemez(self):
        self.giris(self.editor)
        yanit = self.client.post(f"/panel/qr/{self.qr.pk}/duzenle", {
            "ad": "x", "tur": "kampanya", "hedef_turu": "harici", "hedef_url": "https://example.com/cop", "durum": "cop",
        })
        self.assertEqual(yanit.status_code, 200)  # geçersiz seçim: form hatası
        self.qr.refresh_from_db()
        self.assertEqual(self.qr.durum, "aktif")
        self.cope_at(self.qr)
        yanit = self.client.get(f"/panel/qr/{self.qr.pk}/duzenle")
        self.assertRedirects(yanit, f"/panel/qr/{self.qr.pk}/")

    def test_cop_kodlarina_durum_islemi_uygulanmaz(self):
        self.cope_at(self.qr)
        self.client.post("/panel/qr/toplu", {"islem": "aktif", "secili": [self.qr.pk]})
        self.client.post("/panel/qr/toplu", {"islem": "arsivle", "secili": [self.qr.pk], "onay": "1"})
        self.qr.refresh_from_db()
        self.assertEqual((self.qr.durum, self.qr.onceki_durum), ("cop", "aktif"))

    def test_kalici_silme_ucu_yok_ve_delete_engelli(self):
        from django.urls import NoReverseMatch, get_resolver

        adlar = {p.name for p in get_resolver("qr.panel_urls").url_patterns if getattr(p, "name", None)}
        self.assertFalse({a for a in adlar if "sil" in a and "gorsel" not in a})
        self.assertFalse({a for a in adlar if "kalici" in a or "bosalt" in a})
        with self.assertRaises(NoReverseMatch):
            reverse("panel:cop_sil")
        self.giris(self.yonetici)
        for yol in ("/panel/cop-kutusu/sil", "/panel/cop-kutusu/bosalt", f"/panel/qr/{self.qr.pk}/sil"):
            self.assertEqual(self.client.post(yol, {"secili": [self.qr.pk]}).status_code, 404)
        yanit = self.client.get("/panel/cop-kutusu/")
        self.assertNotContains(yanit, "Kalıcı olarak sil")
        self.cope_at(self.qr)
        with self.assertRaises(ValidationError):
            self.qr.delete()
        with self.assertRaises(ValidationError):
            QRKod.objects.get(pk=self.qr.pk).delete()
        self.assertTrue(QRKod.objects.filter(pk=self.qr.pk).exists())

    def test_ice_aktarma_cop_kodunu_guncellemez(self):
        from qr import ice_aktarma
        from qr.models import IceAktarma

        self.cope_at(self.qr)
        ia = IceAktarma.objects.create(
            kullanici=self.editor, dosya_adi="a.csv",
            satirlar=[{"no": 2, "hucreler": ["COP-01", "https://example.com/yeni"]}],
            eslesme={"basliklar": ["kod", "hedef_url"], "alanlar": ["kod", "hedef_url"]},
        )
        analiz = ice_aktarma.analiz_et(ia)
        self.assertEqual(len(analiz.hatali), 1)
        self.assertIn("çöp kutusunda", analiz.hatali[0].hatalar[0][1])
        self.assertEqual(analiz.guncellenecek, [])

    def test_pano_cop_sayisi(self):
        self.cope_at(self.qr)
        self.assertEqual(self.client.get("/panel/").context["sayilar"]["cop"], 1)


class CopApiTesti(TemelTest):
    def setUp(self):
        super().setUp()
        _, self.duz = ApiAnahtari.uret(self.editor, "t")
        self.qr = self.qr_yap("API-C1", "https://example.com/api-cop")
        self.qr.durum = "cop"
        self.qr.onceki_durum = "aktif"
        self.qr.cope_atilma = timezone.now()
        self.qr.save()
        self.arsiv = self.qr_yap("API-A1", "https://example.com/api-arsiv", durum="arsiv")

    def api(self, **p):
        return self.client.get("/api/v1/", p, HTTP_AUTHORIZATION=f"Bearer {self.duz}")

    def test_expand_url_stats_ve_stats_durumu_gosterir(self):
        for kod, durum in (("API-C1", "cop"), ("API-A1", "arsiv")):
            self.assertEqual(self.api(action="expand", shorturl=kod).json()["durum"], durum)
            self.assertEqual(self.api(action="url-stats", shorturl=kod).json()["durum"], durum)
        v = self.api(action="stats", filter="last", limit="10").json()
        durumlar = {l["keyword"]: l["durum"] for l in v["links"]}
        self.assertEqual(durumlar["API-C1"], "cop")
        self.assertEqual(durumlar["API-A1"], "arsiv")
